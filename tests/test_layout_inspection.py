from hashlib import sha256
from io import BytesIO
import json
import unittest

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from src.layout_structure import inspect_layout, map_document, paragraph_text, unsupported_features


def document_bytes(doc):
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def hyperlink(paragraph, target, text):
    rid = paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), rid)
    run = OxmlElement('w:r')
    node = OxmlElement('w:t')
    node.text = text
    run.append(node)
    link.append(run)
    paragraph._p.append(link)
    return link


class LayoutInspectionTests(unittest.TestCase):
    def test_interleaved_block_order_and_nested_cells(self):
        doc = Document()
        doc.add_paragraph('First')
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Left'
        table.cell(0, 1).text = 'Right'
        nested = table.cell(0, 0).add_table(rows=1, cols=1)
        nested.cell(0, 0).text = 'Nested'
        doc.add_paragraph('Last')
        nodes, containers = map_document(doc)
        self.assertEqual(['body/b0', 'body/b1', 'body/b1/r0/c0/b0',
                          'body/b1/r0/c0/b1', 'body/b1/r0/c0/b1/r0/c0/b0',
                          'body/b1/r0/c0/b2', 'body/b1/r0/c1/b0', 'body/b2'], list(nodes))
        self.assertEqual(['body', 'body/b1/r0/c0', 'body/b1/r0/c0/b1/r0/c0',
                          'body/b1/r0/c1'], list(containers))
        _, report = inspect_layout(document_bytes(doc), {})
        self.assertEqual(['body/b0', 'body/b1', 'body/b2'], report['containers'][0]['block_ids'])
        self.assertIn('nested_tables', report['blocks'][1]['unsupported'])
        self.assertIn('nested_tables', report['blocks'][4]['unsupported'])

    def test_merged_cells_are_not_reported_twice(self):
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = 'A'
        table.cell(0, 1).text = 'B'
        table.cell(0, 0).merge(table.cell(0, 1))
        table.cell(1, 0).text = 'C'
        nodes, containers = map_document(doc)
        self.assertEqual(3, len(containers) - 1)
        texts = [paragraph_text(n) for n in nodes.values() if n.tag == qn('w:p')]
        self.assertEqual(1, texts.count('A'))
        self.assertEqual(1, texts.count('B'))
        self.assertIn('merged_cells', unsupported_features(table._tbl))
        self.assertIn('merged_cells', unsupported_features(table.cell(0, 0)._tc.p_lst[0]))

    def test_linked_headers_and_footers_are_deduplicated(self):
        doc = Document()
        doc.add_paragraph('Body')
        doc.sections[0].header.paragraphs[0].text = 'Shared header'
        doc.sections[0].footer.paragraphs[0].text = 'Shared footer'
        doc.add_section(WD_SECTION_START.NEW_PAGE)
        doc.add_section(WD_SECTION_START.NEW_PAGE)
        nodes, containers = map_document(doc)
        self.assertEqual(1, sum(identifier.startswith('header') for identifier in containers))
        self.assertEqual(1, sum(identifier.startswith('footer') for identifier in containers))
        self.assertEqual('Shared header', paragraph_text(nodes['header0/b0']))
        self.assertEqual('Shared footer', paragraph_text(nodes['footer0/b0']))

    def test_inventory_does_not_create_header_footer_parts(self):
        doc = Document()
        doc.add_paragraph('Body')
        part_count = len(doc.part.package.parts)
        nodes, containers = map_document(doc)
        self.assertEqual(['body'], list(containers))
        self.assertEqual(part_count, len(doc.part.package.parts))
        self.assertEqual(['body/b0'], list(nodes))

    def test_first_and_even_header_references_are_inventoried(self):
        doc = Document()
        section = doc.sections[0]
        section.header.paragraphs[0].text = 'Main'
        section.first_page_header.paragraphs[0].text = 'First'
        section.even_page_header.paragraphs[0].text = 'Even'
        nodes, _ = map_document(doc)
        self.assertEqual(['Main', 'First', 'Even'],
                         [paragraph_text(node) for name, node in nodes.items() if name.startswith('header')])

    def test_actual_run_order_links_and_direct_emphasis(self):
        doc = Document()
        p = doc.add_paragraph()
        p.add_run('Backend ').bold = True
        hyperlink(p, 'https://example.test/api', 'API')
        p.add_run(' engineer').italic = True
        _, report = inspect_layout(document_bytes(doc), {})
        block = report['blocks'][0]
        self.assertEqual('Backend API engineer', block['text'])
        self.assertEqual(['Backend ', 'API', ' engineer'], [r['text'] for r in block['runs']])
        self.assertTrue(block['runs'][0]['bold'])
        self.assertTrue(block['runs'][2]['italic'])
        self.assertEqual('https://example.test/api', block['runs'][1]['hyperlink']['target'])
        self.assertTrue(block['runs'][1]['hyperlink']['external'])
        self.assertFalse(block['unsupported'])
        json.dumps(report)

    def test_header_hyperlink_resolves_its_own_relationship_part(self):
        doc = Document()
        hyperlink(doc.sections[0].header.paragraphs[0], 'https://example.test/contact', 'Contact')
        _, report = inspect_layout(document_bytes(doc), {})
        block = next(b for b in report['blocks'] if b['id'] == 'header0/b0')
        self.assertEqual('https://example.test/contact', block['runs'][0]['hyperlink']['target'])

    def test_numbering_indents_and_break_properties_are_reported(self):
        doc = Document()
        p = doc.add_paragraph('Item', style='List Bullet')
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.page_break_before = True
        props = p._p.get_or_add_pPr()
        numbering = OxmlElement('w:numPr')
        num_id = OxmlElement('w:numId')
        num_id.set(qn('w:val'), '42')
        numbering.append(num_id)
        props.append(numbering)
        p.add_run().add_break(WD_BREAK.PAGE)
        _, report = inspect_layout(document_bytes(doc), {})
        block = report['blocks'][0]
        properties = {prop['name']: prop for prop in block['properties']}
        self.assertIn('keepNext', properties)
        self.assertIn('pageBreakBefore', properties)
        self.assertEqual('42', properties['numPr']['children'][0]['attributes']['val'])
        self.assertEqual([{'type': 'page'}], block['runs'][-1]['breaks'])
        self.assertEqual('Item\n', block['text'])

    def test_unsupported_fields_changes_and_drawings_are_reported(self):
        doc = Document()
        p = doc.add_paragraph('Text')
        field = OxmlElement('w:fldSimple')
        field.set(qn('w:instr'), 'PAGE')
        p._p.append(field)
        change = OxmlElement('w:ins')
        run = OxmlElement('w:r')
        text = OxmlElement('w:t')
        text.text = ' inserted'
        run.append(text)
        change.append(run)
        p._p.append(change)
        p.runs[0]._r.append(OxmlElement('w:drawing'))
        _, report = inspect_layout(document_bytes(doc), {})
        self.assertEqual(['drawings', 'fields', 'tracked_changes'], report['blocks'][0]['unsupported'])
        self.assertEqual('Text inserted', report['blocks'][0]['text'])

    def test_report_is_revision_bound_and_rejects_parameters(self):
        doc = Document()
        doc.add_paragraph('Original')
        data = document_bytes(doc)
        changed, report = inspect_layout(data, {})
        self.assertIsNone(changed)
        self.assertEqual(1, report['schema_version'])
        self.assertEqual(sha256(data).hexdigest(), report['workspace_sha256'])
        self.assertEqual('Original', Document(BytesIO(data)).paragraphs[0].text)
        with self.assertRaises(ValueError):
            inspect_layout(data, {'unused': True})

    def test_unaddressable_wrapped_blocks_are_flagged_on_container(self):
        doc = Document()
        doc.add_paragraph('Visible')
        wrapper = OxmlElement('w:sdt')
        contents = OxmlElement('w:sdtContent')
        paragraph = OxmlElement('w:p')
        contents.append(paragraph)
        wrapper.append(contents)
        doc.element.body.insert(0, wrapper)
        _, report = inspect_layout(document_bytes(doc), {})
        self.assertEqual(['content_controls', 'unaddressable_blocks'],
                         report['containers'][0]['unsupported'])
        self.assertEqual(['sdt'], report['containers'][0]['unaddressable_elements'])

    def test_cell_width_and_table_grid_are_reported(self):
        doc = Document()
        doc.add_table(rows=1, cols=2)
        _, report = inspect_layout(document_bytes(doc), {})
        table = report['blocks'][0]
        self.assertEqual(['gridCol', 'gridCol'], [p['name'] for p in table['grid']])
        cell = report['containers'][1]
        self.assertIn('tcW', [p['name'] for p in cell['properties']])


if __name__ == '__main__':
    unittest.main()
