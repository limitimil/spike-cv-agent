import copy
from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document

from src import audit, control
from src.artifacts import Artifacts
from src.storage import lock, read_json, write_json
from src.workflow import Workflow
from src.server import make_server
from src.rendering import DocxToPdfRenderer, RendererError


class FakeRenderer(DocxToPdfRenderer):
    def __init__(self):
        self.version = '1'
        self.calls = 0

    def render(self, source, output):
        self.calls += 1
        Path(output).write_bytes(b'%PDF-1.4\nTEST DOUBLE ONLY\n')

    def validate_environment(self):
        pass

    def metadata(self):
        return {'type': 'test', 'platform': 'test', 'version': self.version}


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        source = self.root / 'input.docx'
        doc = Document()
        doc.add_paragraph('Experience')
        p = doc.add_paragraph()
        p.add_run('Engineer\u00a0').bold = True
        p.add_run('2026')
        doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Skills'
        doc.save(source)
        self.renderer = FakeRenderer()
        self.w = Workflow(self.root, self.renderer)
        self.design = json.loads(Path('design.json').read_text())
        self.design['paragraph_styles'] = {'0': 'Section', '1': 'Role'}
        self.w.init(source, self.design)

    def events(self):
        return [json.loads(line) for line in (self.root / 'audit/events.jsonl').read_text(encoding='utf-8').splitlines()]

    def test_complete_resume_is_idempotent(self):
        self.w.run()
        first = copy.deepcopy(self.w.state['preview'])
        Workflow(self.root, self.renderer).run()
        self.assertEqual(first, read_json(self.root / 'manifest.json')['preview'])
        self.assertEqual(1, self.renderer.calls)
        self.assertEqual(1, sum(e['operation'] == 'content.normalize' for e in self.events()))
        doc = Document(self.w.store.resolve('resume.layout.preview'))
        self.assertTrue(doc.paragraphs[1].runs[0].bold)
        self.assertEqual('Resume.Role', doc.paragraphs[1].style.name)
        self.assertEqual('Engineer 2026', doc.paragraphs[1].text)

    def test_pause_after_atomic_operation_and_resume(self):
        original = self.w.execute
        def execute(*args, **kwargs):
            original(*args, **kwargs)
            control.write(self.root, 'pause')
        with patch.object(self.w, 'execute', side_effect=execute):
            self.w.run()
        self.assertEqual('paused', self.w.state['phase'])
        self.assertEqual(1, self.w.state['next_step'])
        self.assertIsNotNone(self.w.state['preview'])
        control.write(self.root, 'continue')
        resumed = Workflow(self.root, self.renderer)
        resumed.run()
        self.assertEqual('awaiting_review', resumed.state['phase'])

    def test_abort_preserves_preview(self):
        self.w.checkpoint()
        first = copy.deepcopy(self.w.state['preview'])
        control.write(self.root, 'abort')
        self.w.run()
        self.assertEqual(first, self.w.state['preview'])
        self.assertEqual('aborted', self.w.state['phase'])

    def test_stop_after_checkpoint(self):
        control.write(self.root, 'stop_after_checkpoint')
        self.w.run()
        self.assertEqual('stopped', self.w.state['phase'])
        self.assertEqual(1, self.w.state['generation'])

    def test_render_failure_resumes_without_repeating_edits(self):
        with patch.object(self.renderer, 'render', side_effect=RuntimeError('renderer unavailable')):
            with self.assertRaises(RuntimeError):
                self.w.run()
        self.assertEqual('error', self.w.state['phase'])
        self.assertEqual(3, self.w.state['next_step'])
        self.assertNotIn('preview', self.w.state)
        Workflow(self.root, self.renderer).run()
        self.assertEqual(1, sum(e['operation'] == 'layout.apply' for e in self.events()))

    def test_render_failure_preserves_preview_and_debug_checkpoint(self):
        self.w.run()
        preview = copy.deepcopy(self.w.state['preview'])
        pdf = self.w.store.resolve('resume.render.preview').read_bytes()
        self.w.feedback(preview['build_id'], 'Accepted', True)
        self.w.publish(preview['build_id'])
        published = copy.deepcopy(self.w.state['handles']['resume.render.published'])
        self.renderer.version = '2'
        with patch.object(self.renderer, 'render', side_effect=RendererError('DOCUMENT_OPEN_FAILED', 'Cannot open')):
            with self.assertRaises(RendererError):
                self.w.checkpoint()
        self.assertEqual(preview, self.w.state['preview'])
        self.assertEqual(pdf, self.w.store.resolve('resume.render.preview').read_bytes())
        self.assertEqual(published, self.w.state['handles']['resume.render.published'])
        self.assertTrue(self.w.store.resolve('resume.layout.checkpoint').exists())
        event = self.events()[-1]
        self.assertEqual('render.docx_to_pdf', event['operation'])
        self.assertEqual('failed', event['status'])
        self.assertEqual('DOCUMENT_OPEN_FAILED', event['error_code'])
        self.assertEqual(self.w.state['handles']['resume.layout.checkpoint']['build_id'], event['build_id'])

    def test_environment_failure_is_audited_with_checkpoint(self):
        with patch.object(self.renderer, 'validate_environment', side_effect=RendererError('RENDERER_UNAVAILABLE', 'Missing Word')):
            with self.assertRaises(RendererError):
                self.w.checkpoint()
        self.assertNotIn('preview', self.w.state)
        self.assertTrue(self.w.store.resolve('resume.layout.checkpoint').exists())
        self.assertEqual('RENDERER_UNAVAILABLE', self.events()[-1]['error_code'])

    def test_unsupported_platform_fails_before_editing(self):
        self.w.renderer = None
        with patch('src.rendering.factory.platform.system', return_value='Linux'):
            with self.assertRaises(RendererError) as error:
                self.w.run()
        self.assertEqual('UNSUPPORTED_PLATFORM', error.exception.code)
        self.assertEqual(0, self.w.state['next_step'])
        self.assertTrue(self.w.store.resolve('resume.layout.checkpoint').exists())
        self.assertFalse(any(e['operation'] == 'content.normalize' for e in self.events()))

    def test_render_success_uses_shared_build_and_metadata(self):
        self.w.run()
        handles = self.w.state['handles']
        self.assertEqual(handles['resume.layout.preview']['build_id'], handles['resume.render.preview']['build_id'])
        event = next(e for e in self.events() if e['operation'] == 'render.docx_to_pdf')
        self.assertEqual('success', event['status'])
        self.assertEqual(self.renderer.metadata(), event['renderer_metadata'])

    def test_pending_transaction_recovers_once(self):
        state = copy.deepcopy(self.w.state)
        state['phase'] = 'paused'
        ev = audit.event('test.recovery', state['run_id'], reason='Crash recovery test')
        write_json(self.root / '.transaction.json', {'state': state, 'event': ev})
        audit.append(self.root, ev)  # Crash between audit append and manifest replace.
        recovered = Workflow(self.root, self.renderer)
        recovered.store.recover()
        self.assertEqual('paused', recovered.state['phase'])
        self.assertEqual(1, sum(e['event_id'] == ev['event_id'] for e in self.events()))

    def test_torn_audit_tail_recovers(self):
        with (self.root / 'audit/events.jsonl').open('ab') as f:
            f.write(b'{"event_id":"incomplete')
        ev = audit.event('test.recovery', self.w.state['run_id'])
        audit.append(self.root, ev)
        audit.append(self.root, ev)
        lines = (self.root / 'audit/events.jsonl').read_text().splitlines()
        self.assertEqual(ev['event_id'], json.loads(lines[-1])['event_id'])

    def test_feedback_and_publication_gates(self):
        self.w.run()
        build = self.w.state['preview']['build_id']
        with self.assertRaises(ValueError):
            self.w.feedback('wrong-build', 'Too spacious')
        with self.assertRaises(ValueError):
            self.w.publish(build)
        self.w.feedback(build, 'Reviewed every page', True)
        self.w.publish(build)
        self.assertEqual(build, self.w.state['handles']['resume.layout.published']['build_id'])

    def test_immutable_files_and_integrity(self):
        self.w.run()
        entry = copy.deepcopy(self.w.state['handles']['resume.layout.preview'])
        data = (self.root / entry['path']).read_bytes()
        self.w.execute('content.normalize', {}, 'Test another revision')
        self.assertEqual(data, (self.root / entry['path']).read_bytes())
        with self.assertRaises(FileExistsError):
            self.w.store.create(b'changed', '.docx', entry['build_id'])
        (self.root / entry['path']).write_bytes(b'corrupt')
        with self.assertRaises(ValueError):
            self.w.store.resolve('resume.layout.preview')

    def test_writer_lock(self):
        with lock(self.root):
            with self.assertRaises(RuntimeError):
                with lock(self.root):
                    pass

    def test_renderer_identity_change_creates_new_build(self):
        self.w.run()
        first = self.w.state['preview']['build_id']
        self.renderer.version = '2'
        self.w.checkpoint()
        self.assertNotEqual(first, self.w.state['preview']['build_id'])
        self.assertEqual(2, self.renderer.calls)

    def test_control_during_final_render_is_applied(self):
        render = self.renderer.render
        def request_pause(source, output):
            result = render(source, output)
            control.write(self.root, 'pause')
            return result
        with patch.object(self.renderer, 'render', side_effect=request_pause):
            self.w.run()
        self.assertEqual('paused', self.w.state['phase'])

    def test_new_preview_clears_review(self):
        self.w.run()
        self.w.feedback(self.w.state['preview']['build_id'], 'Looks good', True)
        design = dict(self.design, **{'body.size_pt': 11})
        self.w.execute('layout.apply', {'design': design}, 'Increase legibility')
        with self.assertRaises(ValueError):
            self.w.publish(self.w.state['preview']['build_id'])
        self.w.checkpoint()
        self.assertIsNone(self.w.state['review'])

    def test_invalid_design_does_not_change_manifest(self):
        before = (self.root / 'manifest.json').read_bytes()
        design = dict(self.design, **{'body.size_pt': -1})
        with self.assertRaises(ValueError):
            self.w.execute('layout.apply', {'design': design}, 'Invalid request')
        self.assertEqual(before, (self.root / 'manifest.json').read_bytes())

    def test_monitor_exposes_only_published_checkpoints(self):
        import threading
        from urllib.request import urlopen
        from urllib.error import HTTPError
        self.w.run()
        first = copy.deepcopy(self.w.state['preview'])
        server = make_server(self.root, 0)
        worker = threading.Thread(target=server.serve_forever, daemon=True)
        worker.start()
        base = 'http://127.0.0.1:' + str(server.server_port)
        try:
            with urlopen(base + '/preview/status.json') as response:
                self.assertEqual(first, json.load(response)['preview'])
            with urlopen(base + '/preview/current.pdf') as response:
                self.assertEqual(base + '/' + first['pdf'], response.url)
                self.assertTrue(response.read().startswith(b'%PDF-'))
            for path in ('/manifest.json', '/workspace/resume-working.docx',
                         '/' + self.w.state['handles']['resume.source']['path'], '/../../design.json'):
                with self.assertRaises(HTTPError) as error:
                    urlopen(base + path)
                self.assertEqual(404, error.exception.code)
            self.renderer.version = '2'
            self.w.checkpoint()
            with urlopen(base + '/' + first['pdf']) as response:
                self.assertEqual(200, response.status)
        finally:
            server.shutdown()
            server.server_close()
            worker.join(timeout=2)

    def test_browser_controls_require_session_token_and_same_origin(self):
        import re
        import threading
        from urllib.request import urlopen, Request
        from urllib.error import HTTPError
        server = make_server(self.root, 0)
        worker = threading.Thread(target=server.serve_forever, daemon=True)
        worker.start()
        base = 'http://127.0.0.1:' + str(server.server_port)
        try:
            with urlopen(base) as response:
                html = response.read().decode()
            token = re.search("'X-Resume-Token':'([^']+)'", html).group(1)
            def post(path, body, headers=None):
                return urlopen(Request(base + path, data=json.dumps(body).encode(),
                                       headers=headers or {}, method='POST'))
            with self.assertRaises(HTTPError) as error:
                post('/api/control', {'state': 'pause'})
            self.assertEqual(403, error.exception.code)
            with self.assertRaises(HTTPError) as error:
                post('/api/control', {'state': 'pause'}, {'X-Resume-Token': token, 'Origin': 'http://evil.example'})
            self.assertEqual(403, error.exception.code)
            headers = {'X-Resume-Token': token, 'Origin': base}
            with post('/api/control', {'state': 'pause'}, headers) as response:
                self.assertEqual(200, response.status)
            self.assertEqual('pause', control.read(self.root))
            with patch('src.server.subprocess.Popen') as launch:
                launch.return_value.poll.return_value = None
                with post('/api/run', {}, headers) as response:
                    self.assertEqual(200, response.status)
                self.assertEqual('continue', control.read(self.root))
                self.assertEqual('resume', launch.call_args.args[0][-1])
                with self.assertRaises(HTTPError):
                    post('/api/run', {}, headers)
                self.assertEqual(1, launch.call_count)
        finally:
            server.shutdown()
            server.server_close()
            worker.join(timeout=2)


if __name__ == '__main__':
    unittest.main()
