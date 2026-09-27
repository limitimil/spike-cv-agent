from hashlib import sha256
from io import BytesIO
import unittest

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from src.layout_composition import compose_layout, validate_composition
from src.layout_structure import inspect_layout


THEME = {
    'ink': '#173E4C',
    'teal': '#2F6F7E',
    'accent': '#0F6B66',
    'surface': '#EFF5F6',
    'highlight': '#DDEFEA',
    'text': '#232A2E',
    'muted': '#52666D',
    'paper': '#FFFFFF',
}


def document_bytes(doc):
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def specification(data, **changes):
    spec = {'schema_version': 1, 'expected_workspace_sha256': sha256(data).hexdigest(),
            'expected_inspection_sha256': 'a' * 64}
    spec.update(changes)
    return {'specification': spec}


def hyperlink(paragraph, text, target='https://example.test/backend'):
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    value = OxmlElement('w:t')
    value.text = text
    run.append(value)
    link.append(run)
    paragraph._p.append(link)


class LayoutVisualHierarchyTests(unittest.TestCase):
    def compose(self, doc, **changes):
        data = document_bytes(doc)
        return compose_layout(data, specification(data, **changes))

    def test_theme_is_complete_contrast_checked_and_audited(self):
        doc = Document()
        doc.add_paragraph('Skills')
        _, report = self.compose(doc, theme=THEME, paragraph_rules=[{
            'block': 'body/b0', 'color_role': 'ink', 'shading_role': 'surface'}])
        self.assertEqual(7, len(report['theme_contrast']))
        self.assertTrue(all(item['ratio'] >= item['minimum'] for item in report['theme_contrast']))
        self.assertEqual('surface', report['visual_rules'][0]['shading_role'])

        missing = dict(THEME)
        missing.pop('muted')
        with self.assertRaises(ValueError):
            validate_composition(specification(document_bytes(doc), theme=missing))
        unsafe = dict(THEME, muted='#BBBBBB')
        with self.assertRaises(ValueError):
            validate_composition(specification(document_bytes(doc), theme=unsafe))

    def test_paragraph_color_shading_border_alignment_and_tabs_are_applied(self):
        doc = Document()
        for text in ('Skills', 'Centered', 'Right', 'Justified'):
            doc.add_paragraph(text)
        output, _ = self.compose(doc, theme=THEME, paragraph_rules=[
            {'block': 'body/b0', 'alignment': 'left', 'color_role': 'ink',
             'shading_role': 'surface',
             'borders': {'left': {'color_role': 'accent', 'style': 'single',
                                  'width_pt': 2.5, 'space_pt': 4}},
             'tab_stops': [{'position_inches': 2.0, 'alignment': 'left'},
                           {'position_inches': 5.5, 'alignment': 'right', 'leader': 'dots'}]},
            {'block': 'body/b1', 'alignment': 'center'},
            {'block': 'body/b2', 'alignment': 'right'},
            {'block': 'body/b3', 'alignment': 'justify'},
        ])
        result = Document(BytesIO(output))
        self.assertEqual(WD_ALIGN_PARAGRAPH.CENTER, result.paragraphs[1].alignment)
        self.assertEqual(WD_ALIGN_PARAGRAPH.RIGHT, result.paragraphs[2].alignment)
        self.assertEqual(WD_ALIGN_PARAGRAPH.JUSTIFY, result.paragraphs[3].alignment)
        ppr = result.paragraphs[0]._p.pPr
        self.assertEqual('EFF5F6', ppr.find(qn('w:shd')).get(qn('w:fill')))
        left = ppr.find('./' + qn('w:pBdr') + '/' + qn('w:left'))
        self.assertEqual(('single', '20', '4', '0F6B66'),
                         (left.get(qn('w:val')), left.get(qn('w:sz')),
                          left.get(qn('w:space')), left.get(qn('w:color'))))
        tabs = ppr.findall('./' + qn('w:tabs') + '/' + qn('w:tab'))
        self.assertEqual([('left', '2880'), ('right', '7920')],
                         [(tab.get(qn('w:val')), tab.get(qn('w:pos'))) for tab in tabs])
        self.assertEqual('173E4C', result.paragraphs[0].runs[0].font.color.rgb.__str__())

    def test_tab_compaction_preserves_hyperlink_and_mixed_formatting(self):
        doc = Document()
        company = doc.add_paragraph('Example Inc.')
        company.runs[0].bold = True
        date = doc.add_paragraph()
        hyperlink(date, '2022–Present')
        output, report = self.compose(
            doc, theme=THEME,
            paragraph_rules=[{'block': 'body/b0', 'tab_stops': [
                {'position_inches': 5.5, 'alignment': 'right'}]}],
            compact=[{'blocks': ['body/b0', 'body/b1'], 'separator': '\t'}])
        result = Document(BytesIO(output))
        self.assertEqual(1, len(result.paragraphs))
        self.assertEqual('Example Inc.\t2022–Present', result.paragraphs[0].text)
        self.assertTrue(result.paragraphs[0].runs[0].bold)
        _, inspection = inspect_layout(output, {})
        linked = next(run for run in inspection['blocks'][0]['runs'] if 'hyperlink' in run)
        self.assertEqual('https://example.test/backend', linked['hyperlink']['target'])
        self.assertEqual('tab', report['whitespace'][0]['inserted_token'])

    def test_phrase_visual_emphasis_preserves_existing_properties(self):
        doc = Document()
        paragraph = doc.add_paragraph()
        run = paragraph.add_run('Backend APIs at scale')
        run.italic = True
        output, report = self.compose(doc, theme=THEME, highlights=[{
            'block': 'body/b0', 'text': 'Backend APIs', 'expected_matches': 1,
            'color_role': 'teal', 'shading_role': 'highlight', 'bold': True}])
        result = Document(BytesIO(output)).paragraphs[0]
        emphasized = result.runs[0]
        self.assertEqual('Backend APIs', emphasized.text)
        self.assertTrue(emphasized.bold)
        self.assertTrue(emphasized.italic)
        self.assertEqual('2F6F7E', str(emphasized.font.color.rgb))
        self.assertEqual('DDEFEA', emphasized._r.rPr.find(qn('w:shd')).get(qn('w:fill')))
        self.assertEqual(1, report['highlights'][0]['matched'])

    def test_invalid_tabs_fail_before_output(self):
        doc = Document()
        doc.add_paragraph('Company')
        doc.add_paragraph('Date')
        data = document_bytes(doc)
        cases = [
            {'theme': THEME, 'paragraph_rules': [{'block': 'body/b0', 'tab_stops': [
                {'position_inches': 2, 'alignment': 'left'},
                {'position_inches': 1, 'alignment': 'right'}]}]},
            {'compact': [{'blocks': ['body/b0', 'body/b1'], 'separator': '\t'}]},
            {'theme': THEME, 'paragraph_rules': [{'block': 'body/b0', 'tab_stops': [
                {'position_inches': 9, 'alignment': 'right'}]}]},
        ]
        for changes in cases:
            with self.subTest(changes=changes), self.assertRaises(ValueError):
                compose_layout(data, specification(data, **changes))
        self.assertEqual(sha256(data).hexdigest(), sha256(data).hexdigest())

    def test_legacy_spec_behavior_remains_unchanged(self):
        doc = Document()
        doc.add_paragraph('Node.js service')
        output, report = self.compose(doc, highlights=[{
            'block': 'body/b0', 'text': 'Node.js', 'expected_matches': 1}])
        result = Document(BytesIO(output)).paragraphs[0]
        self.assertTrue(result.runs[0].bold)
        self.assertIsNone(result.runs[0].font.color.rgb)
        self.assertEqual([], report['theme_contrast'])


if __name__ == '__main__':
    unittest.main()
