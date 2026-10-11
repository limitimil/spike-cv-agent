import copy
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from src.content_revision import revise_content
from src.layout_structure import paragraph_text
from src.workflow import Workflow
from src import artifacts
from test_workflow import FakeRenderer


def document_bytes(doc):
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def edit(old='old', new='new', block='body/b0', count=1):
    return dict(block=block, old_text=old, new_text=new, expected_matches=count,
                reason='Correct wording', formatting='preserve_uniform')


def parameters(data, edits):
    return {'specification': dict(schema_version=1, expected_workspace_sha256=sha256(data).hexdigest(),
                                 expected_inspection_sha256='a' * 64, edits=edits)}


class RevisionTests(unittest.TestCase):
    def test_uniform_split_runs_and_multiple_input_offsets(self):
        doc = Document()
        p = doc.add_paragraph(style='List Bullet')
        for text in ('old ', 'wo', 'rd old'):
            p.add_run(text).bold = True
        doc.add_paragraph('unaffected')
        data = document_bytes(doc)
        output, report = revise_content(data, parameters(data, [edit(count=2), edit('word', 'phrase')]))
        result = Document(BytesIO(output))
        self.assertEqual('new phrase new', result.paragraphs[0].text)
        self.assertTrue(all(r.bold for r in result.paragraphs[0].runs))
        self.assertEqual('List Bullet', result.paragraphs[0].style.name)
        self.assertEqual([0, 9], report['edits'][0]['matched_offsets'])
        with ZipFile(BytesIO(data)) as before, ZipFile(BytesIO(output)) as after:
            self.assertEqual(before.namelist(), after.namelist())
            for name in before.namelist():
                if name != 'word/document.xml':
                    self.assertEqual(before.read(name), after.read(name), name)

    def test_table_header_deletion_and_unicode(self):
        doc = Document()
        doc.add_paragraph('body')
        doc.add_table(rows=1, cols=1).cell(0, 0).text = 'old text'
        doc.sections[0].header.paragraphs[0].text = 'old header'
        data = document_bytes(doc)
        output, _ = revise_content(data, parameters(data, [edit('old ', '', 'body/b1/r0/c0/b0'),
                                                            edit('old', '新', 'header0/b0')]))
        result = Document(BytesIO(output))
        self.assertEqual('text', result.tables[0].cell(0, 0).text)
        self.assertEqual('新 header', result.sections[0].header.paragraphs[0].text)

    def test_reject_mixed_formatting_controls_and_wrappers(self):
        for kind in ('mixed', 'hyperlink', 'bookmark', 'field', 'tracked', 'section'):
            with self.subTest(kind=kind):
                doc = Document()
                p = doc.add_paragraph()
                p.add_run('o').bold = True
                p.add_run('ld').bold = kind != 'mixed'
                if kind == 'tab':
                    p.add_run('\t')
                elif kind == 'section':
                    p._p.get_or_add_pPr().append(OxmlElement('w:sectPr'))
                elif kind != 'mixed':
                    tag = {'hyperlink': 'hyperlink', 'bookmark': 'bookmarkStart',
                           'field': 'fldSimple', 'tracked': 'ins'}[kind]
                    p._p.append(OxmlElement('w:' + tag))
                data = document_bytes(doc)
                with self.assertRaises(ValueError):
                    revise_content(data, parameters(data, [edit()]))

    def test_unchanged_link_and_numbering_are_preserved(self):
        doc = Document()
        doc.add_paragraph('old')
        p = doc.add_paragraph('numbered', style='List Number')
        link = OxmlElement('w:hyperlink')
        link.set(qn('w:anchor'), 'bookmark')
        run = OxmlElement('w:r')
        text = OxmlElement('w:t')
        text.text = 'link'
        run.append(text)
        link.append(run)
        p._p.append(link)
        data = document_bytes(doc)
        output, _ = revise_content(data, parameters(data, [edit()]))
        self.assertEqual(p._p.xml, Document(BytesIO(output)).paragraphs[1]._p.xml)

    def test_text_on_both_sides_of_controls_preserves_control_xml(self):
        for tag in ('tab', 'br', 'cr'):
            with self.subTest(tag=tag):
                doc = Document()
                p = doc.add_paragraph('old')
                control = OxmlElement('w:' + tag)
                if tag == 'br':
                    control.set(qn('w:type'), 'page')
                p.runs[0]._r.append(control)
                p.add_run('tail')
                before = control.xml
                data = document_bytes(doc)
                output, report = revise_content(data, parameters(data, [edit(), edit('tail', 'longer tail')]))
                changed = Document(BytesIO(output)).paragraphs[0]
                self.assertEqual(before, changed._p.find('.//' + qn('w:' + tag)).xml)
                self.assertEqual('new' + ('\t' if tag == 'tab' else '\n') + 'longer tail', paragraph_text(changed._p))
                self.assertTrue(report['preservation']['control_tokens_unchanged'])
                for bad in (edit('old\ttail', 'merged'), edit('old', 'new\t'), edit('old\ntail', '')):
                    with self.assertRaises(ValueError):
                        revise_content(data, parameters(data, [bad]))

    def test_cross_paragraph_bookmark_interior_is_rejected(self):
        doc = Document()
        first = doc.add_paragraph('before')
        marker = OxmlElement('w:bookmarkStart')
        marker.set(qn('w:id'), '1')
        marker.set(qn('w:name'), 'range')
        first._p.append(marker)
        doc.add_paragraph('old')
        last = doc.add_paragraph('after')
        marker = OxmlElement('w:bookmarkEnd')
        marker.set(qn('w:id'), '1')
        last._p.append(marker)
        data = document_bytes(doc)
        with self.assertRaises(ValueError):
            revise_content(data, parameters(data, [edit(block='body/b1')]))

    def test_proofing_and_empty_goback_are_preserved_with_adjusted_offset(self):
        doc = Document()
        p = doc.add_paragraph('old')
        for kind in ('spellStart', 'spellEnd', 'gramStart', 'gramEnd'):
            marker = OxmlElement('w:proofErr')
            marker.set(qn('w:type'), kind)
            p._p.append(marker)
        start = OxmlElement('w:bookmarkStart')
        start.set(qn('w:id'), '0')
        start.set(qn('w:name'), '_GoBack')
        end = OxmlElement('w:bookmarkEnd')
        end.set(qn('w:id'), '0')
        p._p.append(start)
        p._p.append(end)
        p.add_run(' tail')
        data = document_bytes(doc)
        output, report = revise_content(data, parameters(data, [edit(new='longer'), edit('tail', 'ending')]))
        changed = Document(BytesIO(output)).paragraphs[0]
        self.assertEqual('longer ending', changed.text)
        self.assertEqual(etree.tostring(start), etree.tostring(changed._p.find(qn('w:bookmarkStart'))))
        metadata = report['preserved_metadata']['body/b0']
        self.assertEqual(3, metadata['before']['bookmarks'][0]['offset'])
        self.assertEqual(6, metadata['after']['bookmarks'][0]['offset'])
        with self.assertRaisesRegex(ValueError, 'crosses'):
            revise_content(data, parameters(data, [edit('old tail', 'whole')]))

    def test_invalid_metadata_rejected(self):
        for kind in ('proof_type', 'unpaired', 'named', 'nonempty'):
            with self.subTest(kind=kind):
                doc = Document()
                p = doc.add_paragraph('old')
                if kind == 'proof_type':
                    marker = OxmlElement('w:proofErr')
                    marker.set(qn('w:type'), 'unknown')
                    p._p.append(marker)
                else:
                    start = OxmlElement('w:bookmarkStart')
                    start.set(qn('w:id'), '0')
                    start.set(qn('w:name'), 'other' if kind == 'named' else '_GoBack')
                    p._p.append(start)
                    if kind == 'nonempty':
                        p.add_run('inside')
                    if kind != 'unpaired':
                        end = OxmlElement('w:bookmarkEnd')
                        end.set(qn('w:id'), '0')
                        p._p.append(end)
                data = document_bytes(doc)
                with self.assertRaises(ValueError):
                    revise_content(data, parameters(data, [edit()]))

    def test_invalid_specs_and_overlap(self):
        doc = Document()
        doc.add_paragraph('old old')
        data = document_bytes(doc)
        bad_edits = [[edit()], [edit(block='body/b99')], [edit(count=2), edit(count=2)],
                     [edit(new='bad\ntext', count=2)], [dict(edit(count=2), formatting='guess')],
                     [dict(edit(count=2), expected_matches=True)], [edit('old old', 'old old')]]
        for edits in bad_edits:
            with self.subTest(edits=edits), self.assertRaises(ValueError):
                revise_content(data, parameters(data, edits))
        params = parameters(data, [edit(count=2)])
        params['specification']['expected_workspace_sha256'] = '0' * 64
        with self.assertRaises(ValueError):
            revise_content(data, params)


class RevisionWorkflowTests(unittest.TestCase):
    def test_revision_allowed_at_every_cursor_without_advancement(self):
        for cursor in range(4):
            root = self.root / str(cursor)
            w = Workflow(root, self.renderer)
            w.init(self.root / 'source.docx', self.workflow.state['design'])
            steps = [('source.inspect', {}), ('content.normalize', {}),
                     ('layout.apply', {'design': w.state['design']})]
            for name, args in steps[:cursor]:
                w.execute(name, args, 'Setup legacy cursor', advance=True)
            w.execute('layout.inspect', {}, 'Inspect wording baseline')
            params = parameters(w.store.resolve('resume.layout.workspace').read_bytes(), [edit()])
            params['specification']['expected_inspection_sha256'] = w.state['handles']['resume.layout.inspection']['sha256']
            w.execute('content.revise', params, 'Revise before layout')
            self.assertEqual(cursor, w.state['next_step'])
            self.assertEqual('new', Document(w.store.resolve('resume.layout.workspace')).paragraphs[0].text)

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        source = self.root / 'source.docx'
        doc = Document()
        doc.add_paragraph('old')
        doc.save(source)
        self.renderer = FakeRenderer()
        self.workflow = Workflow(self.root, self.renderer)
        design = json.loads(Path('design.json').read_text(encoding='utf-8'))
        self.workflow.init(source, design)
        self.workflow.run()
        build = self.workflow.state['preview']['build_id']
        self.workflow.feedback(build, 'Explicit test acceptance', True)
        self.workflow.publish(build)
        self.workflow.execute('layout.inspect', {}, 'Inspect before wording edits')
        self.params = parameters(self.workflow.store.resolve('resume.layout.workspace').read_bytes(), [edit()])
        self.params['specification']['expected_inspection_sha256'] = self.workflow.state['handles']['resume.layout.inspection']['sha256']

    def execute(self):
        self.workflow.execute('content.revise', self.params, 'Correct grammar')

    def test_commit_report_invalidate_review_retain_preview_and_render_new(self):
        before = copy.deepcopy(self.workflow.state)
        self.execute()
        state = self.workflow.state
        self.assertIsNone(state['review'])
        self.assertNotIn('resume.layout.inspection', state['handles'])
        self.assertEqual(before['preview'], state['preview'])
        self.assertEqual(before['next_step'], state['next_step'])
        self.assertEqual(before['design'], state['design'])
        for handle in ('resume.source', 'resume.layout.published', 'resume.render.published'):
            self.assertEqual(before['handles'][handle], state['handles'][handle])
        report = json.loads(self.workflow.store.resolve('resume.content.revision').read_text())
        self.assertEqual(state['handles']['resume.layout.workspace']['sha256'], report['output_sha256'])
        self.assertEqual(before['handles']['resume.source'], report['source'])
        with self.assertRaises(ValueError):
            self.workflow.publish(before['preview']['build_id'])
        self.workflow.checkpoint()
        self.assertEqual(2, self.renderer.calls)
        self.assertNotEqual(before['preview']['build_id'], self.workflow.state['preview']['build_id'])
        with self.assertRaises(ValueError):
            self.execute()

    def test_failures_do_not_change_committed_or_in_memory_state(self):
        for field in ('expected_workspace_sha256', 'expected_inspection_sha256'):
            original = self.params['specification'][field]
            self.params['specification'][field] = '0' * 64
            before = copy.deepcopy(self.workflow.state)
            disk = (self.root / 'manifest.json').read_bytes()
            with self.assertRaises(ValueError):
                self.execute()
            self.assertEqual(before, self.workflow.state)
            self.assertEqual(disk, (self.root / 'manifest.json').read_bytes())
            self.params['specification'][field] = original
        with self.assertRaises(ValueError):
            self.workflow.execute('content.revise', self.params, 'No cursor advancement', advance=True)
        before = copy.deepcopy(self.workflow.state)
        create = self.workflow.store.create
        def fail_docx(data, suffix, build=None):
            if suffix == '.docx':
                raise OSError('Interrupted second artifact staging')
            return create(data, suffix, build)
        with patch.object(self.workflow.store, 'create', side_effect=fail_docx), self.assertRaises(OSError):
            self.execute()
        self.assertEqual(before, self.workflow.state)
        self.assertEqual(before, Workflow(self.root, self.renderer).state)

    def test_recovery_commits_report_and_revision_once(self):
        write = artifacts.write_json
        def crash(path, value):
            if Path(path).name == 'manifest.json':
                raise OSError('Interrupted manifest commit')
            return write(path, value)
        with patch('src.artifacts.write_json', side_effect=crash), self.assertRaises(OSError):
            self.execute()
        recovered = Workflow(self.root, self.renderer)
        again = Workflow(self.root, self.renderer)
        self.assertEqual(recovered.state, again.state)
        report = json.loads(recovered.store.resolve('resume.content.revision').read_text())
        self.assertEqual(report['output_sha256'], recovered.state['handles']['resume.layout.workspace']['sha256'])
        events = [json.loads(line) for line in (self.root / 'audit/events.jsonl').read_text().splitlines()]
        self.assertEqual(1, sum(e['operation'] == 'content.revise' for e in events))

    def test_audit_interruption_recovers_both_outputs(self):
        with patch('src.audit.append', side_effect=OSError('Interrupted audit')), self.assertRaises(OSError):
            self.execute()
        recovered = Workflow(self.root, self.renderer)
        recovered.store.resolve('resume.content.revision')
        doc = Document(recovered.store.resolve('resume.layout.workspace'))
        self.assertEqual('new', doc.paragraphs[0].text)
        self.assertIsNone(recovered.state['review'])

    def test_stale_report_integrity_and_invalid_batch_preserve_state(self):
        before = copy.deepcopy(self.workflow.state)
        self.params['specification']['edits'].append(edit('missing', 'changed'))
        with self.assertRaises(ValueError):
            self.execute()
        self.assertEqual(before, self.workflow.state)
        self.params['specification']['edits'].pop()
        report = self.workflow.store.resolve('resume.layout.inspection')
        report.write_bytes(b'{}')
        with self.assertRaisesRegex(ValueError, 'integrity'):
            self.execute()
        self.assertEqual(before, self.workflow.state)
