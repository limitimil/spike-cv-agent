"""Composition must preserve the existing journal, cursor, and review contracts."""
import copy
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document

from src import artifacts
from src.operations import REGISTRY
from src.storage import lock, read_json
from src.workflow import Workflow
from test_workflow import FakeRenderer


class LayoutWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.source = self.root / 'input.docx'
        document = Document()
        document.add_paragraph('Experience')
        paragraph = document.add_paragraph()
        paragraph.add_run('Engineer\u00a0').italic = True
        paragraph.add_run('API delivery')
        document.add_table(rows=1, cols=1).cell(0, 0).text = 'Skills'
        document.add_paragraph('About Me')
        document.save(self.source)
        self.design = json.loads(Path('design.json').read_text(encoding='utf-8'))
        self.renderer = FakeRenderer()
        self.workflow = Workflow(self.root, self.renderer)
        self.workflow.init(self.source, self.design)

    def execute(self, name, params=None, advance=False):
        with lock(self.root):
            return self.workflow.execute(name, params or {},
                                         'Put backend delivery before the experience heading',
                                         advance=advance)

    def run_pipeline(self):
        with lock(self.root):
            self.workflow.run()

    def events(self):
        path = self.root / 'audit/events.jsonl'
        return [json.loads(line) for line in path.read_text(encoding='utf-8').splitlines()]

    def inspect(self):
        self.execute('layout.inspect')
        return read_json(self.workflow.store.resolve('resume.layout.inspection'))

    def specification(self):
        report = self.inspect()
        container = next(item for item in report['containers'] if item['id'] == 'body')
        ordered = list(container['block_ids'])
        ordered[0], ordered[1] = ordered[1], ordered[0]
        return {
            'schema_version': 1,
            'expected_workspace_sha256': self.workflow.state['handles']['resume.layout.workspace']['sha256'],
            'expected_inspection_sha256': self.workflow.state['handles']['resume.layout.inspection']['sha256'],
            'orders': [{'container': 'body', 'blocks': ordered}],
            'typography': {'font': 'Arial', 'body_size_pt': 10.5, 'margin_inches': 0.6},
            'paragraph_rules': [{'block': 'body/b0', 'size_pt': 12, 'bold': True,
                                 'space_before_pt': 8, 'space_after_pt': 3,
                                 'line_spacing': 1, 'keep_with_next': True,
                                 'keep_together': True, 'page_break_before': False,
                                 'widow_control': True, 'indent_inches': 0,
                                 'first_line_indent_inches': 0}],
            'highlights': [{'block': 'body/b1', 'text': 'API', 'expected_matches': 1}],
        }

    def accepted_and_published(self):
        self.run_pipeline()
        build = self.workflow.state['preview']['build_id']
        with lock(self.root):
            self.workflow.feedback(build, 'Reviewed this exact test build', True)
            self.workflow.publish(build)
        return build

    def snapshot(self):
        """Also detect immutable staging created by an invalid specification."""
        return {
            'manifest': (self.root / 'manifest.json').read_bytes(),
            'events': (self.root / 'audit/events.jsonl').read_bytes(),
            'working': (self.root / 'workspace/resume-working.docx').read_bytes(),
            'state': copy.deepcopy(self.workflow.state),
            'artifacts': {path.name: path.read_bytes()
                          for path in (self.root / 'artifacts').iterdir()},
        }

    def assert_rejected_unchanged(self, specification):
        before = self.snapshot()
        with self.assertRaises(ValueError):
            self.execute('layout.compose', {'specification': specification})
        self.assertEqual(before, self.snapshot())

    def test_inspection_preserves_accepted_published_workspace(self):
        self.accepted_and_published()
        before = copy.deepcopy(self.workflow.state)
        working = self.workflow.store.resolve('resume.layout.workspace').read_bytes()
        report = self.inspect()
        self.assertEqual(1, report['schema_version'])
        self.assertEqual(before['handles']['resume.layout.workspace']['sha256'], report['workspace_sha256'])
        role = next(block for block in report['blocks'] if block['id'] == 'body/b1')
        self.assertEqual('Engineer API delivery', role['text'])
        self.assertIn('\u00a0', Document(self.workflow.store.resolve('resume.source')).paragraphs[1].text)
        self.assertEqual(working, self.workflow.store.resolve('resume.layout.workspace').read_bytes())
        for key in ('review', 'preview', 'next_step', 'design', 'phase', 'rendered'):
            self.assertEqual(before[key], self.workflow.state[key], key)
        for handle, entry in before['handles'].items():
            self.assertEqual(entry, self.workflow.state['handles'][handle], handle)
        self.assertEqual(1, self.renderer.calls)
        self.assertEqual('layout.inspect', self.events()[-1]['operation'])

    def test_stale_workspace_and_inspection_are_rejected_atomically(self):
        self.accepted_and_published()
        spec = self.specification()
        wrong_inspection = dict(spec, expected_inspection_sha256='0' * 64)
        self.assert_rejected_unchanged(wrong_inspection)
        newer = dict(self.design, **{'body.size_pt': 11})
        self.execute('layout.apply', {'design': newer})
        self.assert_rejected_unchanged(spec)

    def test_invalid_permutations_highlights_and_unknown_keys_do_not_commit(self):
        self.accepted_and_published()
        spec = self.specification()
        invalid = []
        omitted = copy.deepcopy(spec)
        omitted['orders'][0]['blocks'].pop()
        invalid.append(omitted)
        duplicated = copy.deepcopy(spec)
        duplicated['orders'][0]['blocks'][0] = duplicated['orders'][0]['blocks'][1]
        invalid.append(duplicated)
        missing_phrase = copy.deepcopy(spec)
        missing_phrase['highlights'][0]['text'] = 'Never present'
        invalid.append(missing_phrase)
        invalid.append(dict(spec, unapproved_option=True))
        invalid_size = copy.deepcopy(spec)
        invalid_size['typography']['body_size_pt'] = -1
        invalid.append(invalid_size)
        for bad in invalid:
            with self.subTest(specification=bad):
                self.assert_rejected_unchanged(bad)

    def test_compose_clears_acceptance_and_renders_new_build_without_overwriting_artifacts(self):
        old_build = self.accepted_and_published()
        spec = self.specification()
        old_state = copy.deepcopy(self.workflow.state)
        old_bytes = {handle: self.workflow.store.resolve(handle).read_bytes()
                     for handle in old_state['handles']}
        self.execute('layout.compose', {'specification': spec})
        state = self.workflow.state
        self.assertIsNone(state['review'])
        self.assertEqual(old_state['design'], state['design'])
        self.assertEqual(old_state['next_step'], state['next_step'])
        self.assertEqual(old_state['preview'], state['preview'])
        self.assertNotEqual(old_state['handles']['resume.layout.workspace']['sha256'],
                            state['handles']['resume.layout.workspace']['sha256'])
        for handle, entry in old_state['handles'].items():
            self.assertEqual(old_bytes[handle], (self.root / entry['path']).read_bytes(), handle)
        for handle in ('resume.layout.published', 'resume.render.published'):
            self.assertEqual(old_state['handles'][handle], state['handles'][handle])
        with self.assertRaises(ValueError):
            self.workflow.publish(old_build)
        document = Document(self.workflow.store.resolve('resume.layout.workspace'))
        self.assertEqual(['Engineer API delivery', 'Experience', 'About Me'],
                         [paragraph.text for paragraph in document.paragraphs])
        self.assertEqual('Skills', document.tables[0].cell(0, 0).text)
        self.assertTrue(document.paragraphs[0].runs[0].italic)
        self.assertTrue(any('API' in run.text and run.bold for run in document.paragraphs[0].runs))
        with lock(self.root):
            self.workflow.checkpoint()
        new_build = self.workflow.state['preview']['build_id']
        self.assertNotEqual(old_build, new_build)
        self.assertEqual(2, self.renderer.calls)
        with self.assertRaises(ValueError):
            self.workflow.publish(new_build)
        with lock(self.root):
            self.workflow.feedback(new_build, 'Reviewed the new composed test build', True)
            self.workflow.publish(new_build)
        self.assertEqual(new_build, self.workflow.state['handles']['resume.render.published']['build_id'])
        # The unchanged default resume plan cannot repeat layout.apply over composition.
        composed = self.workflow.state['handles']['resume.layout.workspace']['sha256']
        self.run_pipeline()
        self.assertEqual(composed, self.workflow.state['handles']['resume.layout.workspace']['sha256'])
        self.assertEqual(1, sum(event['operation'] == 'layout.apply' for event in self.events()))
        self.assertEqual(2, self.renderer.calls)

    def test_compose_recovery_commits_one_revision_and_audit_event(self):
        self.run_pipeline()
        spec = self.specification()
        prior_cursor = self.workflow.state['next_step']
        write_json = artifacts.write_json

        def crash_before_manifest(path, value):
            if Path(path).name == 'manifest.json':
                raise OSError('Simulated crash after durable audit append')
            return write_json(path, value)

        with patch('src.artifacts.write_json', side_effect=crash_before_manifest):
            with self.assertRaises(OSError):
                self.execute('layout.compose', {'specification': spec})
        pending = read_json(self.root / '.transaction.json')
        expected = pending['state']['handles']['resume.layout.workspace']
        self.workflow = Workflow(self.root, self.renderer)
        first = copy.deepcopy(self.workflow.state)
        self.workflow = Workflow(self.root, self.renderer)
        self.assertEqual(first, self.workflow.state)
        self.assertEqual(expected, self.workflow.state['handles']['resume.layout.workspace'])
        self.assertEqual(prior_cursor, self.workflow.state['next_step'])
        self.assertEqual(1, sum(event['operation'] == 'layout.compose' for event in self.events()))
        self.assertFalse((self.root / '.transaction.json').exists())
        document = Document(self.workflow.store.resolve('resume.layout.workspace'))
        self.assertEqual('Engineer API delivery', document.paragraphs[0].text)

    def test_staging_crash_leaves_an_orphan_that_recovery_never_selects(self):
        self.run_pipeline()
        spec = self.specification()
        prior = self.snapshot()
        create = self.workflow.store.create
        orphan = []

        def crash_after_staging(data, suffix, build=None):
            entry = create(data, suffix, build)
            orphan.append(entry)
            raise OSError('Simulated crash before journal creation')

        with patch.object(self.workflow.store, 'create', side_effect=crash_after_staging):
            with self.assertRaises(OSError):
                self.execute('layout.compose', {'specification': spec})
        self.assertEqual(1, len(orphan))
        self.assertTrue((self.root / orphan[0]['path']).exists())
        self.assertEqual(prior['manifest'], (self.root / 'manifest.json').read_bytes())
        self.assertEqual(prior['events'], (self.root / 'audit/events.jsonl').read_bytes())
        self.assertFalse((self.root / '.transaction.json').exists())
        self.workflow = Workflow(self.root, self.renderer)
        self.assertEqual(prior['state'], self.workflow.state)
        self.execute('layout.compose', {'specification': spec})
        committed = self.workflow.state['handles']['resume.layout.workspace']
        self.assertNotEqual(orphan[0]['path'], committed['path'])
        self.assertTrue((self.root / orphan[0]['path']).exists())
        self.assertEqual(1, sum(event['operation'] == 'layout.compose' for event in self.events()))

    def test_journal_recovers_composition_when_audit_append_was_interrupted(self):
        self.run_pipeline()
        spec = self.specification()
        with patch('src.audit.append', side_effect=OSError('Simulated audit write failure')):
            with self.assertRaises(OSError):
                self.execute('layout.compose', {'specification': spec})
        pending = read_json(self.root / '.transaction.json')
        self.assertEqual('layout.compose', pending['event']['operation'])
        self.assertFalse(any(event['operation'] == 'layout.compose' for event in self.events()))
        self.workflow = Workflow(self.root, self.renderer)
        self.workflow.store.recover()
        self.assertEqual(pending['state']['handles']['resume.layout.workspace'],
                         self.workflow.state['handles']['resume.layout.workspace'])
        self.assertEqual(1, sum(event['operation'] == 'layout.compose' for event in self.events()))

    def test_early_compose_rejected_at_every_legacy_cursor(self):
        for cursor in range(3):
            with self.subTest(cursor=cursor):
                self.assertEqual(cursor, self.workflow.state['next_step'])
                spec = self.specification()
                self.assert_rejected_unchanged(spec)
                name = ('source.inspect', 'content.normalize', 'layout.apply')[cursor]
                params = {'design': self.design} if name == 'layout.apply' else {}
                self.execute(name, params, advance=True)
        self.assertEqual(3, self.workflow.state['next_step'])
        self.run_pipeline()
        self.assertEqual(1, sum(event['operation'] == 'source.inspect' for event in self.events()))
        self.assertEqual(1, sum(event['operation'] == 'content.normalize' for event in self.events()))
        self.assertEqual(1, sum(event['operation'] == 'layout.apply' for event in self.events()))
        self.assertFalse(any(event['operation'] == 'layout.compose' for event in self.events()))

    def test_pre_feature_manifests_resume_at_each_original_cursor(self):
        steps = [('source.inspect', {}), ('content.normalize', {}),
                 ('layout.apply', {'design': self.design})]
        for cursor in range(4):
            with self.subTest(cursor=cursor):
                root = self.root / ('legacy-' + str(cursor))
                renderer = FakeRenderer()
                workflow = Workflow(root, renderer)
                with lock(root):
                    workflow.init(self.source, self.design)
                    for name, params in steps[:cursor]:
                        workflow.execute(name, params, 'Legacy pipeline step', advance=True)
                manifest = read_json(root / 'manifest.json')
                self.assertEqual(cursor, manifest['next_step'])
                self.assertNotIn('resume.layout.inspection', manifest['handles'])
                resumed = Workflow(root, renderer)
                with lock(root):
                    resumed.run()
                events = [json.loads(line) for line in
                          (root / 'audit/events.jsonl').read_text(encoding='utf-8').splitlines()]
                self.assertEqual(3, resumed.state['next_step'])
                self.assertEqual('awaiting_review', resumed.state['phase'])
                self.assertEqual(self.design, resumed.state['design'])
                self.assertEqual(1, renderer.calls)
                for name, _ in steps:
                    self.assertEqual(1, sum(event['operation'] == name for event in events), name)
                self.assertFalse(any(event['operation'] in ('layout.inspect', 'layout.compose')
                                     for event in events))

    def test_original_operation_shapes_and_design_validation_remain_unchanged(self):
        source = self.workflow.store.resolve('resume.source').read_bytes()
        output, report = REGISTRY['source.inspect'].execute(source, {})
        self.assertIsNone(output)
        self.assertEqual({'paragraphs', 'sections'}, set(report))
        self.assertEqual({'index', 'text', 'style'}, set(report['paragraphs'][0]))
        self.assertEqual(('resume.source',), REGISTRY['source.inspect'].inputs)
        self.assertEqual(('resume.source.inspection',), REGISTRY['source.inspect'].outputs)
        self.assertEqual(('design',), REGISTRY['layout.apply'].parameters)
        REGISTRY['layout.apply'].validate({'design': self.design})
        with self.assertRaises(ValueError):
            REGISTRY['layout.apply'].validate({'design': dict(self.design, orders=[])})
        self.assertEqual(('resume.layout.workspace',), REGISTRY['layout.inspect'].inputs)
        self.assertEqual(('resume.layout.inspection',), REGISTRY['layout.inspect'].outputs)
        self.assertEqual(('specification',), REGISTRY['layout.compose'].parameters)


if __name__ == '__main__':
    unittest.main()

