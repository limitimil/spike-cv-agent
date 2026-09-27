from hashlib import sha256
from io import BytesIO
import unittest
from zipfile import ZipFile

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from src.layout_composition import compose_layout
from src.layout_structure import inspect_layout, map_document


def document_bytes(doc):
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def specification(data, **changes):
    spec = {'schema_version': 1, 'expected_workspace_sha256': sha256(data).hexdigest(),
            'expected_inspection_sha256': 'a' * 64}
    spec.update(changes)
    return {'specification': spec}


def hyperlink(p, text, target='https://example.test/backend'):
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), p.part.relate_to(target, RT.HYPERLINK, is_external=True))
    r = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = text
    r.append(t)
    link.append(r)
    p._p.append(link)
    return r


class LayoutCompositionTests(unittest.TestCase):
    def compose(self, doc, **changes):
        data = document_bytes(doc)
        output, report = compose_layout(data, specification(data, **changes))
        return data, output, report

    def test_highlight_spanning_runs_preserves_unrelated_emphasis(self):
        doc = Document()
        p = doc.add_paragraph()
        p.add_run('Before ').underline = True
        p.add_run('Node.').italic = True
        p.add_run('js ').font.size = Pt(11)
        p.add_run('after').bold = True
        _, output, report = self.compose(doc, highlights=[
            {'block': 'body/b0', 'text': 'Node.js', 'expected_matches': 1}])
        p = Document(BytesIO(output)).paragraphs[0]
        self.assertEqual('Before Node.js after', p.text)
        self.assertTrue(p.runs[0].underline)
        self.assertTrue(p.runs[1].italic)
        self.assertTrue(p.runs[1].bold)
        self.assertTrue(p.runs[2].bold)
        self.assertEqual(11, p.runs[2].font.size.pt)
        self.assertIsNone(p.runs[3].bold)
        self.assertTrue(p.runs[4].bold)
        self.assertEqual(1, report['highlights'][0]['matched'])

    def test_highlight_crossing_hyperlink_keeps_destination_and_context(self):
        doc = Document()
        p = doc.add_paragraph('Backend ')
        r = hyperlink(p, 'API')
        rpr = OxmlElement('w:rPr')
        rpr.append(OxmlElement('w:i'))
        r.insert(0, rpr)
        p.add_run(' design')
        before, output, _ = self.compose(doc, highlights=[
            {'block': 'body/b0', 'text': 'Backend API', 'expected_matches': 1}])
        _, report = inspect_layout(output, {})
        runs = report['blocks'][0]['runs']
        self.assertEqual('Backend API design', report['blocks'][0]['text'])
        self.assertTrue(runs[0]['bold'])
        self.assertTrue(runs[1]['bold'])
        self.assertTrue(runs[1]['italic'])
        self.assertEqual('https://example.test/backend', runs[1]['hyperlink']['target'])
        self.assertNotIn('hyperlink', runs[0])
        self.assertNotIn('hyperlink', runs[2])
        with ZipFile(BytesIO(before)) as original, ZipFile(BytesIO(output)) as changed:
            self.assertEqual(original.read('word/_rels/document.xml.rels'), changed.read('word/_rels/document.xml.rels'))

    def test_repeated_unicode_phrase_requires_explicit_match_count(self):
        doc = Document()
        doc.add_paragraph('資料 API and 資料 API')
        data = document_bytes(doc)
        selector = {'block': 'body/b0', 'text': '資料 API', 'expected_matches': 1}
        with self.assertRaises(ValueError):
            compose_layout(data, specification(data, highlights=[selector]))
        selector['expected_matches'] = 2
        output, _ = compose_layout(data, specification(data, highlights=[selector]))
        _, report = inspect_layout(output, {})
        self.assertEqual([True, None, True], [run['bold'] for run in report['blocks'][0]['runs']])

    def test_interleaved_body_table_permutation_preserves_nested_membership(self):
        doc = Document()
        doc.add_paragraph('Experience')
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = 'Backend Skills'
        doc.add_paragraph('About Me')
        _, output, report = self.compose(doc, orders=[
            {'container': 'body', 'blocks': ['body/b1', 'body/b2', 'body/b0']}])
        result = Document(BytesIO(output))
        nodes, _ = map_document(result)
        self.assertEqual(qn('w:tbl'), nodes['body/b0'].tag)
        self.assertEqual('Backend Skills', result.tables[0].cell(0, 0).text)
        self.assertEqual(['About Me', 'Experience'], [p.text for p in result.paragraphs])
        self.assertTrue(report['preservation']['block_inventory'])

    def test_compact_contact_keeps_link_italic_and_explicit_separator(self):
        doc = Document()
        doc.add_paragraph('Name').runs[0].bold = True
        hyperlink(doc.add_paragraph(), 'Portfolio')
        doc.add_paragraph('Remote').runs[0].italic = True
        _, output, report = self.compose(doc, compact=[
            {'blocks': ['body/b0', 'body/b1', 'body/b2'], 'separator': ' | '}])
        _, inspection = inspect_layout(output, {})
        self.assertEqual(1, len(inspection['blocks']))
        block = inspection['blocks'][0]
        self.assertEqual('Name | Portfolio | Remote', block['text'])
        self.assertTrue(block['runs'][0]['bold'])
        self.assertTrue(block['runs'][-1]['italic'])
        self.assertEqual('https://example.test/backend', block['runs'][2]['hyperlink']['target'])
        self.assertEqual(' | ', report['whitespace'][0]['separator'])

    def test_remove_empty_retains_text_and_untouched_package_parts(self):
        doc = Document()
        doc.add_paragraph('First')
        doc.add_paragraph()
        doc.add_paragraph('Last')
        data, output, _ = self.compose(doc, whitespace=[{'block': 'body/b1', 'action': 'remove_empty'}])
        self.assertEqual(['First', 'Last'], [p.text for p in Document(BytesIO(output)).paragraphs])
        with ZipFile(BytesIO(data)) as before, ZipFile(BytesIO(output)) as after:
            self.assertEqual(before.namelist(), after.namelist())
            for name in before.namelist():
                if name != 'word/document.xml':
                    self.assertEqual(before.read(name), after.read(name), name)

    def test_invalid_orders_and_selectors_fail_without_mutating_input(self):
        doc = Document()
        doc.add_paragraph('One')
        doc.add_paragraph('Two')
        data = document_bytes(doc)
        cases = [
            {'orders': [{'container': 'body', 'blocks': ['body/b0']}]},
            {'orders': [{'container': 'body', 'blocks': ['body/b0', 'body/b0']}]},
            {'orders': [{'container': 'body', 'blocks': ['body/b0', 'body/b99']}]},
            {'paragraph_rules': [{'block': 'body/b99', 'bold': True}]},
            {'paragraph_rules': [{'block': 'body/b0', 'size_pt': float('nan')}]},
            {'highlights': [{'block': 'body/b0', 'text': 'Missing', 'expected_matches': 1}]},
            {'whitespace': [{'block': 'body/b0', 'action': 'remove_empty'}]},
            {'unexpected': True},
        ]
        original = sha256(data).hexdigest()
        for change in cases:
            with self.subTest(change=change), self.assertRaises(ValueError):
                compose_layout(data, specification(data, **change))
            self.assertEqual(original, sha256(data).hexdigest())

    def test_stale_workspace_checksum_is_rejected(self):
        doc = Document()
        doc.add_paragraph('One')
        data = document_bytes(doc)
        params = specification(data)
        params['specification']['expected_workspace_sha256'] = '0' * 64
        with self.assertRaises(ValueError):
            compose_layout(data, params)

    def test_unsupported_affected_fields_and_merged_table_moves_fail(self):
        doc = Document()
        p = doc.add_paragraph('Field')
        p._p.append(OxmlElement('w:fldSimple'))
        data = document_bytes(doc)
        with self.assertRaises(ValueError):
            compose_layout(data, specification(data, paragraph_rules=[{'block': 'body/b0', 'bold': True}]))
        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1))
        doc.add_paragraph('Tail')
        data = document_bytes(doc)
        with self.assertRaises(ValueError):
            compose_layout(data, specification(data, orders=[{'container': 'body', 'blocks': ['body/b1', 'body/b0']}]))

    def test_overlapping_compact_groups_reject_shared_surviving_block(self):
        doc = Document()
        for text in ('First', 'Second', 'Third'):
            doc.add_paragraph(text)
        data = document_bytes(doc)
        with self.assertRaises(ValueError):
            compose_layout(data, specification(data, compact=[
                {'blocks': ['body/b0', 'body/b1'], 'separator': ' '},
                {'blocks': ['body/b0', 'body/b2'], 'separator': ' '},
            ]))

    def test_remove_empty_cannot_delete_cells_only_paragraph_after_table(self):
        doc = Document()
        cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        cell.add_table(rows=1, cols=1)
        cell._tc.remove(cell._tc.p_lst[0])
        data = document_bytes(doc)
        with self.assertRaises(ValueError):
            compose_layout(data, specification(data, whitespace=[
                {'block': 'body/b0/r0/c0/b1', 'action': 'remove_empty'}]))

    def test_remove_empty_cannot_delete_required_trailing_cell_paragraph(self):
        doc = Document()
        cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        cell.paragraphs[0].text = 'Before table'
        cell.add_table(rows=1, cols=1)
        data = document_bytes(doc)
        with self.assertRaises(ValueError):
            compose_layout(data, specification(data, whitespace=[
                {'block': 'body/b0/r0/c0/b2', 'action': 'remove_empty'}]))

    def test_remove_empty_cannot_delete_explicit_numbering_membership(self):
        doc = Document()
        p = doc.add_paragraph()
        numbering = OxmlElement('w:numPr')
        num_id = OxmlElement('w:numId')
        num_id.set(qn('w:val'), '42')
        numbering.append(num_id)
        p._p.get_or_add_pPr().append(numbering)
        data = document_bytes(doc)
        with self.assertRaises(ValueError):
            compose_layout(data, specification(data, whitespace=[
                {'block': 'body/b0', 'action': 'remove_empty'}]))

    def test_compact_cannot_merge_style_inherited_list_members(self):
        doc = Document()
        doc.add_paragraph('First bullet', style='List Bullet')
        doc.add_paragraph('Second bullet', style='List Bullet')
        data = document_bytes(doc)
        with self.assertRaises(ValueError):
            compose_layout(data, specification(data, compact=[
                {'blocks': ['body/b0', 'body/b1'], 'separator': ' '}]))

    def test_affected_run_property_tracked_change_is_rejected(self):
        doc = Document()
        p = doc.add_paragraph('Backend')
        p.runs[0]._r.get_or_add_rPr().append(OxmlElement('w:rPrChange'))
        data = document_bytes(doc)
        with self.assertRaises(ValueError):
            compose_layout(data, specification(data, paragraph_rules=[
                {'block': 'body/b0', 'bold': True}]))

    def test_affected_paragraph_in_merged_table_is_rejected(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Backend'
        table.cell(0, 0).merge(table.cell(0, 1))
        data = document_bytes(doc)
        with self.assertRaises(ValueError):
            compose_layout(data, specification(data, paragraph_rules=[
                {'block': 'body/b0/r0/c0/b0', 'bold': True}]))

    def test_compact_rejects_different_inherited_typography(self):
        doc = Document()
        doc.add_paragraph('Title', style='Heading 1')
        doc.add_paragraph('Plain contact', style='Normal')
        data = document_bytes(doc)
        with self.assertRaises(ValueError):
            compose_layout(data, specification(data, compact=[
                {'blocks': ['body/b0', 'body/b1'], 'separator': ' '}]))

    def test_remove_empty_rejects_style_inherited_numbering(self):
        doc = Document()
        doc.add_paragraph('', style='List Bullet')
        data = document_bytes(doc)
        with self.assertRaises(ValueError):
            compose_layout(data, specification(data, whitespace=[
                {'block': 'body/b0', 'action': 'remove_empty'}]))

    def test_highlight_preserves_non_target_run_properties_and_attributes(self):
        doc = Document()
        p = doc.add_paragraph()
        run = p.add_run('Before API after')
        run.font.small_caps = True
        run.font.subscript = True
        run.italic = True
        run._r.set(qn('w:rsidR'), '00112233')
        _, output, _ = self.compose(doc, highlights=[
            {'block': 'body/b0', 'text': 'API', 'expected_matches': 1}])
        p = Document(BytesIO(output)).paragraphs[0]
        self.assertEqual(['Before ', 'API', ' after'], [r.text for r in p.runs])
        self.assertEqual([None, True, None], [r.bold for r in p.runs])
        for split in p.runs:
            self.assertTrue(split.font.small_caps)
            self.assertTrue(split.font.subscript)
            self.assertTrue(split.italic)
            self.assertEqual('00112233', split._r.get(qn('w:rsidR')))

    def test_compact_differing_paragraph_defaults_need_direct_overrides(self):
        doc = Document()
        for text, size in [('Label', '24'), ('Value', '20')]:
            p = doc.add_paragraph(text)
            defaults = OxmlElement('w:rPr')
            prop = OxmlElement('w:sz')
            prop.set(qn('w:val'), size)
            defaults.append(prop)
            p._p.get_or_add_pPr().append(defaults)
        data = document_bytes(doc)
        with self.assertRaises(ValueError):
            compose_layout(data, specification(data, compact=[
                {'blocks': ['body/b0', 'body/b1'], 'separator': ' '}]))

    def test_compact_differing_defaults_with_complete_overrides_preserves_runs(self):
        doc = Document()
        for text, size in [('Label', '24'), ('Value', '20')]:
            p = doc.add_paragraph(text)
            defaults = OxmlElement('w:rPr')
            prop = OxmlElement('w:sz')
            prop.set(qn('w:val'), size)
            defaults.append(prop)
            p._p.get_or_add_pPr().append(defaults)
        _, output, _ = self.compose(doc, paragraph_rules=[
            {'block': 'body/b0', 'size_pt': 11}, {'block': 'body/b1', 'size_pt': 10.5}],
            compact=[{'blocks': ['body/b0', 'body/b1'], 'separator': ' '}])
        result = Document(BytesIO(output)).paragraphs[0]
        self.assertEqual('Label Value', result.text)
        self.assertEqual(11, result.runs[0].font.size.pt)
        self.assertEqual(10.5, result.runs[-1].font.size.pt)

    def test_compact_rejects_incomplete_overrides_with_unrelated_extra_attributes(self):
        doc = Document()
        for text, east_asia in [('Label', 'DefaultA'), ('Value', 'DefaultB')]:
            p = doc.add_paragraph(text)
            defaults = OxmlElement('w:rPr')
            prop = OxmlElement('w:rFonts')
            prop.set(qn('w:ascii'), 'Same')
            prop.set(qn('w:eastAsia'), east_asia)
            defaults.append(prop)
            p._p.get_or_add_pPr().append(defaults)
            direct = OxmlElement('w:rFonts')
            direct.set(qn('w:ascii'), 'Explicit')
            direct.set(qn('w:hAnsi'), 'Explicit')
            p.runs[0]._r.get_or_add_rPr().append(direct)
        data = document_bytes(doc)
        with self.assertRaises(ValueError):
            compose_layout(data, specification(data, compact=[
                {'blocks': ['body/b0', 'body/b1'], 'separator': ' '}]))

    def test_highlight_preserves_local_bookmark_boundaries(self):
        doc = Document()
        p = doc.add_paragraph('Before ')
        start = OxmlElement('w:bookmarkStart')
        start.set(qn('w:id'), '7')
        start.set(qn('w:name'), '_GoBack')
        p._p.append(start)
        p.add_run('Node.').italic = True
        p.add_run('js')
        end = OxmlElement('w:bookmarkEnd')
        end.set(qn('w:id'), '7')
        p._p.append(end)
        p.add_run(' after')
        _, output, report = self.compose(doc, highlights=[
            {'block': 'body/b0', 'text': 'Node.js', 'expected_matches': 1}])
        result = Document(BytesIO(output)).paragraphs[0]
        self.assertEqual('Before Node.js after', result.text)
        position, boundaries = 0, []
        for child in result._p.iter():
            if child.tag == qn('w:t'):
                position += len(child.text or '')
            elif child.tag in (qn('w:bookmarkStart'), qn('w:bookmarkEnd')):
                boundaries.append((child.tag, dict(child.attrib), position))
        self.assertEqual([(qn('w:bookmarkStart'), {qn('w:id'): '7', qn('w:name'): '_GoBack'}, 7),
                          (qn('w:bookmarkEnd'), {qn('w:id'): '7'}, 14)], boundaries)
        self.assertTrue(result.runs[1].bold)
        self.assertTrue(result.runs[1].italic)
        self.assertTrue(report['preservation']['bookmarks'])

    def test_highlight_keeps_leading_cached_pagebreak_once_and_run_formatting(self):
        doc = Document()
        p = doc.add_paragraph()
        r = p.add_run('API Performance Testing')
        r.italic = True
        r.font.size = Pt(11)
        r._r.insert(1, OxmlElement('w:lastRenderedPageBreak'))
        _, output, _ = self.compose(doc, highlights=[
            {'block': 'body/b0', 'text': 'API Performance Testing', 'expected_matches': 1}])
        result = Document(BytesIO(output)).paragraphs[0]
        self.assertEqual('API Performance Testing', result.text)
        self.assertEqual(1, len(list(result._p.iter(qn('w:lastRenderedPageBreak')))))
        text_runs = [run for run in result.runs if run.text]
        self.assertTrue(all(run.bold and run.italic for run in text_runs))
        self.assertTrue(all(run.font.size.pt == 11 for run in text_runs))

    def test_highlight_spanning_cached_pagebreak_preserves_marker_position(self):
        doc = Document()
        p = doc.add_paragraph()
        r = p.add_run('API Per')
        r.italic = True
        r._r.append(OxmlElement('w:lastRenderedPageBreak'))
        trailing = OxmlElement('w:t')
        trailing.text = 'formance Testing'
        r._r.append(trailing)
        _, output, _ = self.compose(doc, highlights=[
            {'block': 'body/b0', 'text': 'API Performance Testing', 'expected_matches': 1}])
        result = Document(BytesIO(output)).paragraphs[0]
        position, markers = 0, []
        for child in result._p.iter():
            if child.tag == qn('w:t'):
                position += len(child.text or '')
            elif child.tag == qn('w:lastRenderedPageBreak'):
                markers.append(position)
        self.assertEqual([7], markers)
        self.assertEqual('API Performance Testing', result.text)
        self.assertTrue(all(run.bold and run.italic for run in result.runs if run.text))

    def test_reorder_rejects_bookmark_spanning_paragraphs(self):
        doc = Document()
        first = doc.add_paragraph('Start')
        second = doc.add_paragraph('End')
        start = OxmlElement('w:bookmarkStart')
        start.set(qn('w:id'), '3')
        start.set(qn('w:name'), 'Span')
        first._p.insert(0, start)
        end = OxmlElement('w:bookmarkEnd')
        end.set(qn('w:id'), '3')
        second._p.append(end)
        data = document_bytes(doc)
        with self.assertRaises(ValueError):
            compose_layout(data, specification(data, orders=[
                {'container': 'body', 'blocks': ['body/b1', 'body/b0']}]))

    def test_explicit_font_removes_conflicting_theme_references_preserves_italic(self):
        doc = Document()
        p = doc.add_paragraph('Backend typography')
        run = p.runs[0]
        run.italic = True
        fonts = OxmlElement('w:rFonts')
        for attr in ('asciiTheme', 'hAnsiTheme', 'csTheme', 'cstheme'):
            fonts.set(qn('w:' + attr), 'minorHAnsi')
        run._r.get_or_add_rPr().append(fonts)
        _, output, _ = self.compose(doc, typography={'font': 'Arial'})
        result = Document(BytesIO(output)).paragraphs[0].runs[0]
        fonts = result._r.find('./' + qn('w:rPr') + '/' + qn('w:rFonts'))
        for attr in ('asciiTheme', 'hAnsiTheme', 'csTheme', 'cstheme'):
            self.assertNotIn(qn('w:' + attr), fonts.attrib)
        for attr in ('ascii', 'hAnsi', 'cs'):
            self.assertEqual('Arial', fonts.get(qn('w:' + attr)))
        self.assertEqual('Backend typography', result.text)
        self.assertTrue(result.italic)

    def test_explicit_indent_removes_character_overrides_and_keeps_numbering(self):
        doc = Document()
        p = doc.add_paragraph('Backend bullet')
        ppr = p._p.get_or_add_pPr()
        ind = OxmlElement('w:ind')
        for attr, value in [('leftChars', '100'), ('firstLineChars', '200'),
                            ('hangingChars', '300'), ('rightChars', '400')]:
            ind.set(qn('w:' + attr), value)
        ppr.append(ind)
        numbering = OxmlElement('w:numPr')
        num_id = OxmlElement('w:numId')
        num_id.set(qn('w:val'), '42')
        numbering.append(num_id)
        ppr.append(numbering)
        _, output, _ = self.compose(doc, paragraph_rules=[
            {'block': 'body/b0', 'indent_inches': .2, 'first_line_indent_inches': -.2}])
        result = Document(BytesIO(output)).paragraphs[0]
        ind = result._p.find('./' + qn('w:pPr') + '/' + qn('w:ind'))
        for attr in ('leftChars', 'firstLineChars', 'hangingChars'):
            self.assertNotIn(qn('w:' + attr), ind.attrib)
        self.assertEqual('288', ind.get(qn('w:left')))
        self.assertEqual('288', ind.get(qn('w:hanging')))
        self.assertEqual('400', ind.get(qn('w:rightChars')))
        num_id = result._p.find('./' + qn('w:pPr') + '/' + qn('w:numPr') + '/' + qn('w:numId'))
        self.assertEqual('42', num_id.get(qn('w:val')))
        self.assertEqual('Backend bullet', result.text)


if __name__ == '__main__':
    unittest.main()
