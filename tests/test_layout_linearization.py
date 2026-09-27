"""The table-to-body extension preserves content and existing composition limits."""
import copy
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor
from docx.text.run import Run

from src import artifacts
from src.layout_composition import compose_layout, inventory, canonical
from src.layout_structure import inspect_layout, map_document
from src.storage import lock, read_json
from src.workflow import Workflow
from test_workflow import FakeRenderer


def document_bytes(doc):
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def specification(data, **changes):
    spec = {'schema_version': 1, 'expected_workspace_sha256': sha256(data).hexdigest(),
            'expected_inspection_sha256': 'a' * 64}
    spec.update(changes)
    return {'specification': spec}


def source_layout():
    """Horizontal contact/sidebar merges and nested company heading, like source."""
    doc = Document()
    table = doc.add_table(rows=2, cols=3)
    contact = table.cell(0, 0).merge(table.cell(0, 1))
    contact.paragraphs[0].text = 'Name'
    contact.paragraphs[0].runs[0].bold = True
    contact.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string('FFFFFF')
    portfolio = table.cell(0, 2).paragraphs[0]
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), portfolio.part.relate_to(
        'https://example.test/portfolio', RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    properties = OxmlElement('w:rPr')
    properties.append(OxmlElement('w:i'))
    run.append(properties)
    text = OxmlElement('w:t')
    text.text = 'Portfolio'
    run.append(text)
    link.append(run)
    portfolio._p.append(link)
    skills = table.cell(1, 0)
    skills.paragraphs[0].text = 'Skills'
    bullet = skills.add_paragraph('Backend API delivery')
    numbering = OxmlElement('w:numPr')
    level = OxmlElement('w:ilvl')
    level.set(qn('w:val'), '0')
    number = OxmlElement('w:numId')
    number.set(qn('w:val'), '1')
    numbering.extend([level, number])
    bullet._p.get_or_add_pPr().append(numbering)
    experience = table.cell(1, 1).merge(table.cell(1, 2))
    experience.paragraphs[0].text = 'About Me'
    nested = experience.add_table(rows=1, cols=2)
    nested.cell(0, 0).text = 'Company'
    nested.cell(0, 1).text = 'Engineer'
    experience.add_paragraph('Experience')
    doc.add_paragraph('Tail')
    return doc


def table_paragraphs(data, table='body/b0'):
    _, report = inspect_layout(data, {})
    return [block for block in report['blocks']
            if block['type'] == 'paragraph' and block['id'].startswith(table + '/')]


def linearization(data, order=None):
    ids = [block['id'] for block in table_paragraphs(data)]
    return [{'table': 'body/b0', 'paragraph_order': ids if order is None else order}]


class LayoutLinearizationTests(unittest.TestCase):
    def test_horizontal_merges_and_nested_table_become_declared_body_paragraph_order(self):
        data = document_bytes(source_layout())
        blocks = table_paragraphs(data)
        requested = list(reversed([block['id'] for block in blocks]))
        original_nodes, _ = map_document(Document(BytesIO(data)))
        original_tokens = [inventory(original_nodes[block]) for block in requested]
        original_numbering = [canonical(original_nodes[block].find(
            './' + qn('w:pPr') + '/' + qn('w:numPr'))) for block in requested]
        output, result = compose_layout(data, specification(
            data, linearize_tables=linearization(data, requested)))
        document = Document(BytesIO(output))
        nodes, _ = map_document(document)
        self.assertEqual([], document.tables)
        self.assertEqual(len(blocks) + 1, len(document.paragraphs))
        self.assertEqual([inventory(nodes['body/b' + str(i)]) for i in range(len(blocks))],
                         original_tokens)
        self.assertEqual([canonical(nodes['body/b' + str(i)].find(
            './' + qn('w:pPr') + '/' + qn('w:numPr'))) for i in range(len(blocks))],
                         original_numbering)
        self.assertEqual('Tail', document.paragraphs[-1].text)
        self.assertTrue(result['preservation']['text'])
        self.assertTrue(result['preservation']['hyperlinks'])
        self.assertTrue(result['preservation']['numbering'])
        self.assertTrue(result['preservation']['block_inventory'])
        audit = result['linearized_tables'][0]
        self.assertEqual('body/b0', audit['table'])
        self.assertEqual(requested, audit['paragraph_order'])
        self.assertEqual({block: index for index, block in enumerate(requested)}, audit['body_positions'])
        self.assertEqual(2, len(audit['removed_table_shells']))
        for shell in audit['removed_table_shells']:
            self.assertEqual(64, len(shell['xml_sha256']))
        _, report = inspect_layout(output, {})
        linked = next(block for block in report['blocks'] if block.get('text') == 'Portfolio')
        self.assertEqual('https://example.test/portfolio', linked['runs'][0]['hyperlink']['target'])
        self.assertTrue(linked['runs'][0]['italic'])
        with ZipFile(BytesIO(data)) as before, ZipFile(BytesIO(output)) as after:
            self.assertEqual(before.namelist(), after.namelist())
            for name in before.namelist():
                if name != 'word/document.xml':
                    self.assertEqual(before.read(name), after.read(name), name)

    def test_original_ids_remain_targetable_for_rules_highlights_and_contact_compaction(self):
        data = document_bytes(source_layout())
        blocks = table_paragraphs(data)
        by_text = {block['text']: block['id'] for block in blocks if block['text']}
        contacts = [by_text['Name'], by_text['Portfolio']]
        order = contacts + [block['id'] for block in blocks if block['id'] not in contacts]
        output, report = compose_layout(data, specification(
            data, linearize_tables=linearization(data, order),
            typography={'font': 'Arial', 'body_size_pt': 10.5, 'margin_inches': 0.6,
                        'color': '#21343B'},
            paragraph_rules=[{'block': by_text['Skills'], 'size_pt': 13,
                              'bold': True, 'keep_with_next': True}],
            highlights=[{'block': by_text['Backend API delivery'], 'text': 'API',
                         'expected_matches': 1}],
            compact=[{'blocks': contacts, 'separator': ' | '}] ))
        document = Document(BytesIO(output))
        self.assertEqual([], document.tables)
        self.assertEqual('Name | Portfolio', document.paragraphs[0].text)
        _, inspection = inspect_layout(output, {})
        contact = inspection['blocks'][0]
        self.assertTrue(contact['runs'][0]['bold'])
        self.assertTrue(contact['runs'][-1]['italic'])
        self.assertEqual('https://example.test/portfolio', contact['runs'][-1]['hyperlink']['target'])
        heading = next(p for p in document.paragraphs if p.text == 'Skills')
        self.assertEqual(13, heading.runs[0].font.size.pt)
        self.assertTrue(heading.runs[0].bold)
        self.assertTrue(heading.paragraph_format.keep_with_next)
        bullet = next(p for p in document.paragraphs if p.text == 'Backend API delivery')
        self.assertTrue(any(run.text == 'API' and run.bold for run in bullet.runs))
        self.assertEqual('1', bullet._p.find('./' + qn('w:pPr') + '/' + qn('w:numPr')
                                          + '/' + qn('w:numId')).get(qn('w:val')))
        # The inserted separator has no explicit color; every source run has the requested one.
        for element in document.element.iter(qn('w:r')):
            run = Run(element, None)
            if run.text == ' | ':
                continue
            self.assertEqual('21343B', str(run.font.color.rgb))
        self.assertEqual(1, report['highlights'][0]['matched'])
        self.assertEqual('#21343B', report['specification']['typography']['color'])

    def test_missing_duplicate_foreign_and_unknown_paragraph_ids_are_rejected(self):
        data = document_bytes(source_layout())
        order = [block['id'] for block in table_paragraphs(data)]
        cases = [order[:-1], order + [order[0]], ['body/b1'] + order[1:],
                 ['body/b0/r99/c99/b99'] + order[1:], ['body/b0'] + order[1:]]
        before = sha256(data).hexdigest()
        for bad in cases:
            with self.subTest(order=bad), self.assertRaises(ValueError):
                compose_layout(data, specification(data, linearize_tables=linearization(data, bad)))
            self.assertEqual(before, sha256(data).hexdigest())
        with self.assertRaises(ValueError):
            compose_layout(data, specification(data, linearize_tables=linearization(data) * 2))

    def test_stale_specification_cannot_linearize_another_revision(self):
        data = document_bytes(source_layout())
        params = specification(data, linearize_tables=linearization(data))
        params['specification']['expected_workspace_sha256'] = '0' * 64
        with self.assertRaises(ValueError):
            compose_layout(data, params)

    def test_vertical_and_legacy_horizontal_merges_are_rejected(self):
        for merge in ('vMerge', 'hMerge'):
            with self.subTest(merge=merge):
                doc = Document()
                table = doc.add_table(rows=2, cols=1)
                table.cell(0, 0).text = 'One'
                table.cell(1, 0).text = 'Two'
                marker = OxmlElement('w:' + merge)
                marker.set(qn('w:val'), 'restart')
                table.cell(0, 0)._tc.get_or_add_tcPr().append(marker)
                data = document_bytes(doc)
                with self.assertRaises(ValueError):
                    compose_layout(data, specification(data, linearize_tables=linearization(data)))

    def test_fields_drawings_wrapped_content_and_changes_reject_linearization(self):
        for tag in ('fldSimple', 'drawing', 'sdt', 'ins', 'rPrChange'):
            with self.subTest(tag=tag):
                doc = Document()
                paragraph = doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
                paragraph.text = 'Preserve this'
                if tag == 'rPrChange':
                    paragraph.runs[0]._r.get_or_add_rPr().append(OxmlElement('w:' + tag))
                else:
                    paragraph._p.append(OxmlElement('w:' + tag))
                data = document_bytes(doc)
                with self.assertRaises(ValueError):
                    compose_layout(data, specification(data, linearize_tables=linearization(data)))

    def test_only_body_level_tables_can_be_selected(self):
        data = document_bytes(source_layout())
        _, report = inspect_layout(data, {})
        nested = next(block['id'] for block in report['blocks']
                      if block['type'] == 'table' and block['id'] != 'body/b0')
        params = specification(data, linearize_tables=[{
            'table': nested, 'paragraph_order': [p['id'] for p in table_paragraphs(data, nested)]}])
        with self.assertRaises(ValueError):
            compose_layout(data, params)

    def test_old_specifications_keep_merged_nested_restrictions_and_original_color(self):
        data = document_bytes(source_layout())
        output, _ = compose_layout(data, specification(data))
        _, report = inspect_layout(output, {})
        name = next(block for block in report['blocks'] if block.get('text') == 'Name')
        color = next(prop for prop in name['runs'][0]['properties'] if prop['name'] == 'color')
        self.assertEqual('FFFFFF', color['attributes']['val'])
        self.assertEqual(1, len(Document(BytesIO(output)).tables))
        for change in ({'paragraph_rules': [{'block': name['id'], 'bold': True}]},
                       {'orders': [{'container': 'body', 'blocks': ['body/b1', 'body/b0']}]},
                       {'typography': {'font': 'Arial'}}):
            with self.subTest(change=change), self.assertRaises(ValueError):
                compose_layout(data, specification(data, **change))

    def test_color_is_explicit_and_validated(self):
        doc = Document()
        paragraph = doc.add_paragraph('White contact text')
        paragraph.runs[0].font.color.rgb = RGBColor.from_string('FFFFFF')
        data = document_bytes(doc)
        unchanged, _ = compose_layout(data, specification(data, typography={'font': 'Arial'}))
        self.assertEqual('FFFFFF', str(Document(BytesIO(unchanged)).paragraphs[0].runs[0].font.color.rgb))
        output, _ = compose_layout(data, specification(data, typography={'color': '#21343B'}))
        self.assertEqual('21343B', str(Document(BytesIO(output)).paragraphs[0].runs[0].font.color.rgb))
        for invalid in ('white', '#123', '#GGGGGG', '21343B', '', True, None):
            with self.subTest(color=invalid), self.assertRaises(ValueError):
                compose_layout(data, specification(data, typography={'color': invalid}))


class LayoutLinearizationWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        source = self.root / 'input.docx'
        source.write_bytes(document_bytes(source_layout()))
        design = json.loads(Path('design.json').read_text(encoding='utf-8'))
        self.workflow = Workflow(self.root, FakeRenderer())
        with lock(self.root):
            self.workflow.init(source, design)
            self.workflow.run()
            build = self.workflow.state['preview']['build_id']
            self.workflow.feedback(build, 'Reviewed exact pre-linearization test build', True)
            self.workflow.publish(build)
            self.workflow.execute('layout.inspect', {}, 'Inspect structural table layout')
        data = self.workflow.store.resolve('resume.layout.workspace').read_bytes()
        self.params = specification(data, linearize_tables=linearization(data))
        self.params['specification']['expected_inspection_sha256'] = (
            self.workflow.state['handles']['resume.layout.inspection']['sha256'])

    def test_rejection_preserves_committed_state_and_all_artifact_bytes(self):
        invalid = copy.deepcopy(self.params)
        invalid['specification']['linearize_tables'][0]['paragraph_order'].pop()
        manifest = (self.root / 'manifest.json').read_bytes()
        events = (self.root / 'audit/events.jsonl').read_bytes()
        state = copy.deepcopy(self.workflow.state)
        files = {path.name: path.read_bytes() for path in (self.root / 'artifacts').iterdir()}
        with lock(self.root), self.assertRaises(ValueError):
            self.workflow.execute('layout.compose', invalid, 'Invalid incomplete table order')
        self.assertEqual(manifest, (self.root / 'manifest.json').read_bytes())
        self.assertEqual(events, (self.root / 'audit/events.jsonl').read_bytes())
        self.assertEqual(state, self.workflow.state)
        self.assertEqual(files, {path.name: path.read_bytes()
                                 for path in (self.root / 'artifacts').iterdir()})

    def test_linearization_recovers_once_without_reusing_prior_acceptance(self):
        prior = copy.deepcopy(self.workflow.state)
        write_json = artifacts.write_json

        def interrupted_manifest(path, value):
            if Path(path).name == 'manifest.json':
                raise OSError('Interrupted linearization manifest commit')
            return write_json(path, value)

        with patch('src.artifacts.write_json', side_effect=interrupted_manifest):
            with lock(self.root), self.assertRaises(OSError):
                self.workflow.execute('layout.compose', self.params,
                                      'Replace the layout table with a single reading column')
        pending = read_json(self.root / '.transaction.json')
        recovered = Workflow(self.root, self.workflow.renderer)
        recovered.store.recover()
        self.assertEqual(pending['state']['handles']['resume.layout.workspace'],
                         recovered.state['handles']['resume.layout.workspace'])
        self.assertEqual([], Document(recovered.store.resolve('resume.layout.workspace')).tables)
        self.assertIsNone(recovered.state['review'])
        self.assertEqual(prior['preview'], recovered.state['preview'])
        self.assertEqual(prior['next_step'], recovered.state['next_step'])
        self.assertEqual(prior['design'], recovered.state['design'])
        self.assertEqual(prior['handles']['resume.render.published'],
                         recovered.state['handles']['resume.render.published'])
        events = [json.loads(line) for line in
                  (self.root / 'audit/events.jsonl').read_text(encoding='utf-8').splitlines()]
        self.assertEqual(1, sum(event['operation'] == 'layout.compose' for event in events))
        with self.assertRaises(ValueError):
            recovered.publish(prior['preview']['build_id'])


if __name__ == '__main__':
    unittest.main()
