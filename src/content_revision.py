"""Exact editorial patches; no generated wording and no structural edits."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import re
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from .layout_structure import map_document, paragraph_text, unsupported_features, _part_for_element, _run_text


def validate_revision(params):
    if not isinstance(params, dict) or set(params) != {'specification'}:
        raise ValueError('Expected specification')
    spec = params['specification']
    keys = {'schema_version', 'expected_workspace_sha256', 'expected_inspection_sha256', 'edits'}
    if not isinstance(spec, dict) or set(spec) != keys:
        raise ValueError('Invalid revision specification keys')
    if type(spec['schema_version']) is not int or spec['schema_version'] != 1:
        raise ValueError('Expected schema_version 1')
    for key in ('expected_workspace_sha256', 'expected_inspection_sha256'):
        if not isinstance(spec[key], str) or not re.fullmatch('[0-9a-f]{64}', spec[key]):
            raise ValueError('Invalid checksum: ' + key)
    if not isinstance(spec['edits'], list) or not spec['edits']:
        raise ValueError('edits must be a nonempty list')
    for edit in spec['edits']:
        if not isinstance(edit, dict) or set(edit) != {
                'block', 'old_text', 'new_text', 'expected_matches', 'reason', 'formatting'}:
            raise ValueError('Invalid edit keys')
        for key in ('block', 'old_text', 'reason'):
            if not isinstance(edit[key], str) or not edit[key].strip():
                raise ValueError('Nonempty ' + key + ' required')
        if not isinstance(edit['new_text'], str) or edit['new_text'] == edit['old_text']:
            raise ValueError('Replacement must change text')
        for key in ('old_text', 'new_text'):
            if any(ord(c) < 32 or 0xD800 <= ord(c) <= 0xDFFF or ord(c) in (0xFFFE, 0xFFFF)
                   for c in edit[key]):
                raise ValueError('Text cannot contain controls or invalid XML characters')
        if type(edit['expected_matches']) is not int or edit['expected_matches'] < 1:
            raise ValueError('expected_matches must be a positive integer')
        if edit['formatting'] != 'preserve_uniform':
            raise ValueError('formatting must be preserve_uniform')


def _metadata(paragraph):
    """Validate narrowly supported Word metadata and inventory text positions."""
    children = list(paragraph)
    bookmarks, proofing = [], []
    offset = 0
    paired_ends = set()
    ids = set()
    for index, child in enumerate(children):
        if child.tag == qn('w:r'):
            offset += len(_run_text(child))
        elif child.tag == qn('w:proofErr'):
            if (len(child) or set(child.attrib) != {qn('w:type')} or
                    child.get(qn('w:type')) not in ('spellStart', 'spellEnd', 'gramStart', 'gramEnd')):
                raise ValueError('Malformed proofing marker')
            proofing.append(child.get(qn('w:type')))
        elif child.tag == qn('w:bookmarkStart'):
            identifier = child.get(qn('w:id'))
            end = children[index + 1] if index + 1 < len(children) else None
            if (len(child) or child.get(qn('w:name')) != '_GoBack' or not identifier or
                    identifier in ids or set(child.attrib) != {qn('w:id'), qn('w:name')} or
                    end is None or end.tag != qn('w:bookmarkEnd') or len(end) or
                    dict(end.attrib) != {qn('w:id'): identifier}):
                raise ValueError('Only paired adjacent empty _GoBack bookmarks are supported')
            ids.add(identifier)
            paired_ends.add(index + 1)
            bookmarks.append({'id': identifier, 'offset': offset})
        elif child.tag == qn('w:bookmarkEnd') and index not in paired_ends:
            raise ValueError('Unpaired bookmark end')
    return {'proofing': proofing, 'bookmarks': bookmarks}


def _text_nodes(paragraph):
    _metadata(paragraph)
    if unsupported_features(paragraph):
        raise ValueError('Unsupported affected paragraph structure')
    if paragraph.find('.//' + qn('w:sectPr')) is not None:
        raise ValueError('Section boundaries are unsupported')
    result = []
    offset = 0
    for child in paragraph:
        if child.tag in (qn('w:pPr'), qn('w:proofErr'), qn('w:bookmarkStart'), qn('w:bookmarkEnd')):
            continue
        if child.tag != qn('w:r'):
            raise ValueError('Affected paragraphs must contain plain runs only (no links/bookmarks/wrappers)')
        properties = child.find(qn('w:rPr'))
        signature = b'' if properties is None else etree.tostring(properties, method='c14n')
        for node in child:
            if node.tag == qn('w:rPr'):
                continue
            if node.tag == qn('w:lastRenderedPageBreak'):
                continue  # nonprinting cache marker remains exactly where it was
            if node.tag in (qn('w:tab'), qn('w:br'), qn('w:cr')):
                offset += 1  # same token width as paragraph_text; never a text target
                continue
            if node.tag != qn('w:t'):
                raise ValueError('Affected paragraphs contain unsupported controls')
            length = len(node.text or '')
            result.append((offset, offset + length, node, signature))
            offset += length
    return result


def _check_enclosing_ranges(doc, paragraph):
    """An interior paragraph may have no local marker for a spanning range."""
    active = set()
    fields = 0
    for node in _part_for_element(doc, paragraph).element.iter():
        if node is paragraph:
            if active or fields:
                raise ValueError('Paragraph is inside a cross-paragraph range or field')
            return
        if node.tag in (qn('w:bookmarkStart'), qn('w:commentRangeStart')):
            active.add((node.tag, node.get(qn('w:id'))))
        elif node.tag in (qn('w:bookmarkEnd'), qn('w:commentRangeEnd')):
            start = 'w:bookmarkStart' if node.tag == qn('w:bookmarkEnd') else 'w:commentRangeStart'
            active.discard((qn(start), node.get(qn('w:id'))))
        elif node.tag == qn('w:fldChar'):
            kind = node.get(qn('w:fldCharType'))
            if kind == 'begin':
                fields += 1
            elif kind == 'end':
                fields = max(0, fields - 1)


def revise_content(data, params):
    validate_revision(params)
    spec = params['specification']
    if sha256(data).hexdigest() != spec['expected_workspace_sha256']:
        raise ValueError('Stale revision workspace checksum')
    doc = Document(BytesIO(data))
    blocks, _ = map_document(doc)
    plans, details = {}, []
    for edit in spec['edits']:
        identifier = edit['block']
        paragraph = blocks.get(identifier)
        if paragraph is None or paragraph.tag != qn('w:p'):
            raise ValueError('Unknown paragraph: ' + identifier)
        _check_enclosing_ranges(doc, paragraph)
        nodes = _text_nodes(paragraph)
        text = paragraph_text(paragraph)
        matches = [m.start() for m in re.finditer('(?=' + re.escape(edit['old_text']) + ')', text)]
        if len(matches) != edit['expected_matches']:
            raise ValueError('Unexpected text match count: ' + identifier)
        changes = plans.setdefault(identifier, [])
        for start in matches:
            end = start + len(edit['old_text'])
            if any(start < bookmark['offset'] < end for bookmark in _metadata(paragraph)['bookmarks']):
                raise ValueError('Replacement crosses an empty bookmark')
            if any(start < prior_end and end > prior_start for prior_start, prior_end, _ in changes):
                raise ValueError('Overlapping text edits')
            affected = [item for item in nodes if item[0] < end and item[1] > start]
            if sum(min(end, item[1]) - max(start, item[0]) for item in affected) != end - start:
                raise ValueError('Replacement crosses a control token')
            if len({item[3] for item in affected}) != 1:
                raise ValueError('Replacement spans mixed run formatting')
            changes.append((start, end, edit['new_text']))
        details.append(dict(edit, matched_offsets=matches))

    changed_parts, expected_text = {}, {}
    original_parts, preserved_metadata = {}, {}
    for identifier, changes in plans.items():
        paragraph = blocks[identifier]
        part = _part_for_element(doc, paragraph)
        name = str(part.partname).lstrip('/')
        if name not in original_parts:
            original_parts[name] = deepcopy(part.element)
        changed_parts[name] = part
        text = paragraph_text(paragraph)
        metadata = _metadata(paragraph)
        # All offsets refer to the input revision; reverse order keeps them valid.
        for start, end, replacement in sorted(changes, reverse=True):
            affected = [item for item in _text_nodes(paragraph) if item[0] < end and item[1] > start]
            for index, (left, right, node, _) in enumerate(affected):
                current = node.text or ''
                a, b = max(start - left, 0), min(end - left, right - left)
                node.text = current[:a] + (replacement if index == 0 else '') + current[b:]
                node.set(qn('xml:space'), 'preserve')
            text = text[:start] + replacement + text[end:]
        expected_text[identifier] = text
        if paragraph_text(paragraph) != text:
            raise ValueError('Declared text diff validation failed')
        expected_metadata = deepcopy(metadata)
        for bookmark in expected_metadata['bookmarks']:
            bookmark['offset'] += sum(len(replacement) - (end - start)
                                      for start, end, replacement in changes if end <= bookmark['offset'])
        if _metadata(paragraph) != expected_metadata:
            raise ValueError('Metadata position preservation failed')
        preserved_metadata[identifier] = {'before': metadata, 'after': expected_metadata}

    # Only text nodes in declared paragraphs may differ. All other XML, including
    # links, numbering, bookmarks, run properties and section settings, is exact.
    for name, part in changed_parts.items():
        before, after = original_parts[name], deepcopy(part.element)
        before_nodes, after_nodes = list(before.iter()), list(after.iter())
        if len(before_nodes) != len(after_nodes):
            raise ValueError('Unexpected structure change')
        for old, new in zip(before_nodes, after_nodes):
            if old.tag == qn('w:t') and new.tag == old.tag:
                new.text = old.text
                if qn('xml:space') in old.attrib:
                    new.set(qn('xml:space'), old.get(qn('xml:space')))
                else:
                    new.attrib.pop(qn('xml:space'), None)
        if etree.tostring(before) != etree.tostring(after):
            raise ValueError('Unexpected non-text XML change')

    output = BytesIO()
    with ZipFile(BytesIO(data)) as source, ZipFile(output, 'w') as target:
        if len(set(source.namelist())) != len(source.namelist()):
            raise ValueError('Duplicate package members are unsupported')
        target.comment = source.comment
        for item in source.infolist():
            payload = changed_parts[item.filename].blob if item.filename in changed_parts else source.read(item)
            target.writestr(item, payload)
    result = output.getvalue()
    # Validate the saved package and every paragraph, not just the intended edits.
    saved, _ = map_document(Document(BytesIO(result)))
    original, _ = map_document(Document(BytesIO(data)))
    if list(saved) != list(original):
        raise ValueError('Block membership changed')
    for identifier, paragraph in original.items():
        if paragraph.tag == qn('w:p'):
            expected = expected_text.get(identifier, paragraph_text(paragraph))
            if paragraph_text(saved[identifier]) != expected:
                raise ValueError('Saved text differs from declared diff')
    return result, {
        'schema_version': 1, 'specification': spec, 'edits': details,
        'preserved_metadata': preserved_metadata,
        'input_sha256': sha256(data).hexdigest(), 'output_sha256': sha256(result).hexdigest(),
        'preservation': {'declared_diff_only': True, 'non_text_xml_unchanged': True,
                         'control_tokens_unchanged': True,
                         'unaffected_package_members_unchanged': True},
    }
