from dataclasses import dataclass
from io import BytesIO
import math

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches, Pt

from .layout_structure import inspect_layout
from .layout_composition import compose_layout, validate_composition
from .content_revision import revise_content, validate_revision


@dataclass(frozen=True)
class Operation:
    name: str
    inputs: tuple
    outputs: tuple
    parameters: tuple
    validate: object
    execute: object
    audit_serializer: object
    lifecycle: bool = False


REGISTRY = {}


def register(operation):
    if operation.name in REGISTRY:
        raise ValueError('Duplicate operation: ' + operation.name)
    REGISTRY[operation.name] = operation


def paragraphs(doc):
    def walk(container):
        yield from container.paragraphs
        for table in container.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from walk(cell)
    seen = set()
    containers = [doc]
    for section in doc.sections:
        containers.extend([section.header, section.footer])
    for container in containers:
        for paragraph in walk(container):
            if paragraph._p not in seen:
                seen.add(paragraph._p)
                yield paragraph


def validate_design(params):
    d = params['design']
    allowed = {'body.font', 'body.size_pt', 'section.space_before_pt',
               'experience.role_spacing_pt', 'page.margin_inches',
               'pagination.widow_control', 'paragraph_styles'}
    if set(d) - allowed:
        raise ValueError('Unknown design keys: ' + str(set(d) - allowed))
    for key, low, high in [('body.size_pt', 6, 30), ('section.space_before_pt', 0, 72),
                            ('experience.role_spacing_pt', 0, 72), ('page.margin_inches', .2, 2)]:
        value = d.get(key)
        if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or not low <= value <= high:
            raise ValueError('Invalid design value: ' + key)
    if not isinstance(d.get('body.font'), str) or not d['body.font'].strip():
        raise ValueError('body.font must be nonempty')
    if not isinstance(d.get('pagination.widow_control'), bool):
        raise ValueError('pagination.widow_control must be boolean')
    mapping = d.get('paragraph_styles', {})
    if not isinstance(mapping, dict) or any(v not in ('Body', 'Section', 'Company', 'Role', 'Meta', 'Bullet') for v in mapping.values()):
        raise ValueError('Invalid paragraph_styles mapping')


def inspect(data, params):
    doc = Document(BytesIO(data))
    return None, {'paragraphs': [{'index': i, 'text': p.text, 'style': p.style.name}
                               for i, p in enumerate(paragraphs(doc))], 'sections': len(doc.sections)}


def normalize(data, params):
    doc = Document(BytesIO(data))
    changed = 0
    for p in paragraphs(doc):
        for run in p.runs:
            text = run.text.replace('\u00a0', ' ')
            if text != run.text:
                run.text = text
                changed += 1
    output = BytesIO()
    doc.save(output)
    return output.getvalue(), {'normalized_runs': changed}


def layout(data, params):
    doc = Document(BytesIO(data))
    d = params['design']
    for name in ('Body', 'Section', 'Company', 'Role', 'Meta', 'Bullet'):
        full = 'Resume.' + name
        style = doc.styles[full] if full in doc.styles else doc.styles.add_style(full, WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = d['body.font']
        style.font.size = Pt(d['body.size_pt'])
        style.paragraph_format.widow_control = d['pagination.widow_control']
        if name in ('Section', 'Company', 'Role'):
            style.font.bold = True
            style.paragraph_format.keep_with_next = True
        if name == 'Section':
            style.paragraph_format.space_before = Pt(d['section.space_before_pt'])
        if name == 'Role':
            style.paragraph_format.space_before = Pt(d['experience.role_spacing_pt'])
    ps = list(paragraphs(doc))
    mapping = d.get('paragraph_styles', {})
    if any(not k.isdigit() or int(k) >= len(ps) or str(int(k)) != k for k in mapping):
        raise ValueError('paragraph_styles contains an invalid paragraph index; use source.inspect')
    for i, p in enumerate(ps):
        if str(i) in mapping:
            p.style = doc.styles['Resume.' + mapping[str(i)]]
            p.paragraph_format.space_before = None
            p.paragraph_format.keep_with_next = None
        for run in p.runs:
            run.font.name = d['body.font']
            run.font.size = Pt(d['body.size_pt'])
        p.paragraph_format.widow_control = d['pagination.widow_control']
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(d['page.margin_inches'])
        section.left_margin = section.right_margin = Inches(d['page.margin_inches'])
    output = BytesIO()
    doc.save(output)
    return output.getvalue(), {'design': d, 'styled_paragraphs': mapping}


def no_validation(params):
    if params:
        raise ValueError('This operation takes no parameters')


def serialize(before, after, reason):
    return {'before': before, 'after': after, 'reason': reason}


for name, fn, validation, parameters, output in [
    ('source.inspect', inspect, no_validation, (), ('resume.source.inspection',)),
    ('content.normalize', normalize, no_validation, (), ('resume.layout.workspace',)),
    ('layout.apply', layout, validate_design, ('design',), ('resume.layout.workspace',))]:
    register(Operation(name, ('resume.source' if name == 'source.inspect' else 'resume.layout.workspace',),
                       output, parameters, validation, fn, serialize))

register(Operation('layout.inspect', ('resume.layout.workspace',), ('resume.layout.inspection',),
                   (), no_validation, inspect_layout, serialize))
register(Operation('layout.compose', ('resume.layout.workspace', 'resume.layout.inspection'),
                   ('resume.layout.workspace',), ('specification',), validate_composition, compose_layout, serialize))
register(Operation('content.revise', ('resume.layout.workspace', 'resume.layout.inspection'),
                   ('resume.layout.workspace', 'resume.content.revision'),
                   ('specification',), validate_revision, revise_content, serialize))


def lifecycle_validator(required):
    def validate(params):
        if set(params) != set(required):
            raise ValueError('Expected parameters: ' + ', '.join(required))
    return validate


for name, inputs, outputs, keys, fn in [
    ('render.docx_to_pdf', ('resume.layout.workspace',), ('resume.layout.checkpoint', 'resume.render.candidate'), (), lambda w, p: w.render()),
    ('preview.publish', ('resume.render.candidate',), ('resume.layout.preview', 'resume.render.preview'), (), lambda w, p: w.checkpoint()),
    ('visual.review', ('resume.render.preview',), (), ('build', 'message', 'accepted'), lambda w, p: w.feedback(p['build'], p['message'], p['accepted'])),
    ('artifact.publish', ('resume.layout.preview', 'resume.render.preview'), ('resume.layout.published', 'resume.render.published'), ('build',), lambda w, p: w.publish(p['build']))]:
    register(Operation(name, inputs, outputs, keys, lifecycle_validator(keys), fn, serialize, True))
